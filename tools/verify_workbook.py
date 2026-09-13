#!/usr/bin/env python3
"""
tools/verify_workbook.py - structural QA for the Catering Business Manager.

Opens every product workbook with openpyxl and checks:

  * the right tabs exist (per edition), _Data is hidden, tab order
  * the defined names the formulas rely on are present
  * the _Data KPI table is complete (35 rows, no gaps)
  * no cached value is an Excel error (#REF!, #VALUE!, ...)
  * validation / conditional-format / chart counts are sane
  * sheets are protected when a password build is expected
  * EXAMPLE builds contain the demo business; blank builds don't
  * the P&L NET PROFIT formula is GROSS - OPEX (regression check)
  * every internal hyperlink points at a real sheet

Usage:
    python3 tools/verify_workbook.py [file.xlsx ...]     # default: products/
    python3 tools/verify_workbook.py --expect-protect
"""

import glob
import os
import sys
import warnings

warnings.filterwarnings("ignore")

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from catering_tracker import config as C  # noqa: E402

ERRORS = ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!",
          "#NUM!")

EXPECTED_SHEETS = {
    "premium": [C.SHEET_NAMES[k] for k in C.EDITIONS["premium"]],
    "basic": [C.SHEET_NAMES[k] for k in C.EDITIONS["basic"]],
}


def check(path, expect_protect=False):
    import openpyxl

    fails = []
    name = os.path.basename(path)
    edition = "premium" if "PREMIUM" in name else "basic"
    demo = "_EXAMPLE" in name

    wb = openpyxl.load_workbook(path)
    titles = [ws.title for ws in wb.worksheets]

    if titles != EXPECTED_SHEETS[edition]:
        fails.append("tab set/order: %s" % titles)

    hidden = [ws.title for ws in wb.worksheets if ws.sheet_state != "visible"]
    if hidden != [C.SHEET_NAMES["data"]]:
        fails.append("hidden sheets: %s" % hidden)

    need = ["BusinessName", "Currency", "TaxRate", "DefaultMargin",
            "DepositPct", "DueSoonDays", "CalMonth", "CalYear",
            "ReportYear", "EventTypes", "ExpenseCategories",
            "PaymentMethods", "StaffRoles", "EventStatuses",
            "PaymentStatuses", "Tick", "EventList", "ClientsList"]
    if edition == "premium":
        need += ["MenuCategories", "IngredientCategories", "MenuItems",
                 "SuppliersList"]
    missing = [n for n in need if n not in wb.defined_names]
    if missing:
        fails.append("missing defined names: %s" % missing)

    # KPI table complete
    d = wb[C.SHEET_NAMES["data"]]
    kpi_rows = 0
    for row in range(2, 37):
        if d.cell(row=row, column=31).value:
            kpi_rows += 1
    if kpi_rows != 35:
        fails.append("KPI table has %d/35 labels" % kpi_rows)

    # cached error scan (skip the literal '#' header cells)
    wbv = openpyxl.load_workbook(path, data_only=True)
    bad = []
    for ws in wbv.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v in ERRORS:
                    bad.append("%s!%s=%s" % (ws.title, cell.coordinate, v))
    if bad:
        fails.append("cached errors: %s" % bad[:5])

    # feature counts
    ndv = sum(len(ws.data_validations.dataValidation)
              for ws in wb.worksheets)
    ncf = sum(len(ws.conditional_formatting._cf_rules)
              for ws in wb.worksheets)
    nch = sum(len(ws._charts) for ws in wb.worksheets)
    min_dv, min_cf, min_ch = ((60, 60, 8) if edition == "premium"
                              else (40, 8, 4))
    if ndv < min_dv:
        fails.append("only %d validations (<%d)" % (ndv, min_dv))
    if ncf < min_cf:
        fails.append("only %d cond formats (<%d)" % (ncf, min_cf))
    if nch < min_ch:
        fails.append("only %d charts (<%d)" % (nch, min_ch))

    # protection
    if expect_protect:
        unprotected = [ws.title for ws in wb.worksheets
                       if not ws.protection.sheet]
        if unprotected:
            fails.append("unprotected sheets: %s" % unprotected)

    # demo presence
    dash = wbv[C.SHEET_NAMES["dashboard"]]
    hero = dash["B2"].value
    if demo and hero != "Saffron & Sage Catering Co.":
        fails.append("EXAMPLE build hero = %r" % hero)
    if not demo and hero not in ("Your Catering Business", ""):
        fails.append("blank build hero = %r" % hero)

    # P&L NET PROFIT regression (doc bug: must be $D$11-$D$12)
    rep = wb[C.SHEET_NAMES["reports"]] if C.SHEET_NAMES["reports"] in \
        titles else None
    if rep is not None:
        net = rep["D13"].value
        if net != '=IF($D$12="","",$D$11-$D$12)':
            fails.append("P&L NET PROFIT formula = %r" % net)

    # internal links point at real sheets
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.hyperlink is not None and \
                        cell.hyperlink.location is not None:
                    target = cell.hyperlink.location.split("!")[0]
                    target = target.strip("'")
                    if target not in titles:
                        fails.append("bad link %s!%s -> %s"
                                     % (ws.title, cell.coordinate, target))

    ok = not fails
    print("%-4s %-52s %s" % ("PASS" if ok else "FAIL", name,
                             "" if ok else "; ".join(fails)[:400]))
    return ok


def main(argv):
    args = [a for a in argv if not a.startswith("--")]
    expect_protect = "--expect-protect" in argv
    paths = args or sorted(glob.glob("products/Catering_Business_Manager_"
                                     "*.xlsx"))
    if not paths:
        print("no workbooks found (build first)")
        return 1
    print("Verifying %d workbook(s)\n" % len(paths))
    ok = all(check(p, expect_protect) for p in paths)
    print("\n%s" % ("ALL CHECKS PASSED" if ok else "FAILURES FOUND"))
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
