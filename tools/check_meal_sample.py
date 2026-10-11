#!/usr/bin/env python3
"""
Check that the example data in a built workbook hangs together.

    python tools/check_meal_sample.py "output/Meal Planner & Grocery List.xlsx"

The lookups in this workbook all match on names, so the example data is only
as good as its spelling.  This tool walks the registers and reports every
ingredient, recipe, store, unit, category and meal slot that is used on one
tab but missing from another.
"""

from __future__ import annotations

import argparse
import os
import sys
from collections import Counter, defaultdict

try:
    import openpyxl
except ImportError:                              # pragma: no cover
    print("error: openpyxl is required (pip install openpyxl)",
          file=sys.stderr)
    raise SystemExit(2)

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

DEFAULT_OUTPUT = os.path.join("output",
                              "Meal Planner & Grocery List.xlsx")


def column_map(ws, header_row=4):
    """{column key guess: column index} from the header row."""
    headers = {}
    for cell in ws[header_row]:
        if cell.value:
            headers[str(cell.value).strip()] = cell.column
    return headers


def read_column(ws, index, first_row, last_row):
    if index is None:
        return []
    return [ws.cell(row=r, column=index).value
            for r in range(first_row, last_row + 1)]


def non_blank(values):
    return [v for v in values if v not in (None, "")]


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", nargs="?", default=DEFAULT_OUTPUT)
    parser.add_argument("--rows", type=int, default=400,
                        help="how many rows of each register to read")
    args = parser.parse_args(argv)

    if not os.path.exists(args.workbook):
        print(f"error: {args.workbook} does not exist", file=sys.stderr)
        return 2

    wb = openpyxl.load_workbook(args.workbook, data_only=True, read_only=True)
    problems = []
    notes = []

    def sheet(name):
        return wb[name] if name in wb.sheetnames else None

    def col(ws, header, header_row=4):
        return column_map(ws, header_row).get(header)

    def values(name, header, header_row=4, first_row=None):
        ws = sheet(name)
        if ws is None:
            return []
        index = col(ws, header, header_row)
        if index is None:
            return []
        first = first_row or header_row + 1
        last = first + args.rows
        return non_blank(read_column(ws, index, first, last))

    # ---- the reference lists ------------------------------------------------
    lists = sheet("Lists")
    list_sets = {}
    if lists is not None:
        headers = column_map(lists, 4)
        for name, index in headers.items():
            list_sets[name] = set(non_blank(
                read_column(lists, index, 5, 5 + args.rows)))

    ingredients = set(values("Ingredients", "Ingredient"))
    recipes = set(values("Recipes", "Recipe"))
    stores = set(values("Lists", "Stores")) | set(values("Stores", "Store"))
    units = set(values("Lists", "Units"))
    categories = set(values("Lists", "Categories"))
    slots = set(values("Lists", "Meal slots"))
    statuses = set(values("Lists", "Status"))

    recipe_categories = set(values("Lists", "Recipe categories"))
    notes.append(f"{len(ingredients)} ingredients, {len(recipes)} recipes, "
                 f"{len(stores)} stores, {len(units)} units, "
                 f"{len(categories)} aisle categories, "
                 f"{len(recipe_categories)} recipe categories, "
                 f"{len(slots)} meal slots")

    if not ingredients:
        problems.append("the Ingredients tab is empty - run the build with "
                        "the example data (drop --blank)")
    if not recipes:
        problems.append("the Recipes tab is empty - run the build with the "
                        "example data (drop --blank)")

    # ---- duplicates ----------------------------------------------------------
    for label, seq in (("ingredient", values("Ingredients", "Ingredient")),
                       ("recipe", values("Recipes", "Recipe"))):
        dupes = [k for k, v in Counter(seq).items() if v > 1]
        if dupes:
            problems.append(f"duplicate {label} names: {sorted(dupes)[:8]}")

    # ---- names used on one tab must exist on another -------------------------
    checks = [
        ("Recipe Ingredients", "Ingredient", ingredients, "Ingredients"),
        ("Recipe Ingredients", "Recipe", recipes, "Recipes"),
        ("Meal Plan", "Recipe", recipes, "Recipes"),
        ("Meal Plan", "Meal slot", slots, "Lists"),
        ("Meal Plan", "Status", statuses, "Lists"),
        ("Grocery List", "Item", ingredients, "Ingredients"),
        ("Grocery List", "Store", stores, "Lists"),
        ("Grocery List", "Unit", units, "Lists"),
        ("Pantry", "Item", ingredients, "Ingredients"),
        ("Pantry", "Storage", {"Pantry", "Cupboard", "Fridge", "Freezer"},
         "storage"),
        ("Staples", "Staple", ingredients, "Ingredients"),
        ("Leftovers", "From recipe", recipes, "Recipes"),
        ("Food Waste", "Item", ingredients, "Ingredients"),
        ("Meal Prep", "Recipe", recipes, "Recipes"),
        ("Recipe Costs", "Recipe", recipes, "Recipes"),
        ("Rotation", "Recipe", recipes, "Recipes"),
        ("Fridge & Freezer", "From recipe", recipes, "Recipes"),
        ("Nutrition", "Recipe", recipes, "Recipes"),
        ("Nutrition", "Meal slot", slots, "Lists"),
        ("Grocery Spending", "Item", ingredients, "Ingredients"),
        ("Grocery Spending", "Store", stores, "Lists"),
        ("Store Prices", "Item", ingredients, "Ingredients"),
        ("Store Prices", "Store", stores, "Lists"),
        ("Ingredients", "Category", categories, "Lists"),
        ("Ingredients", "Unit", units, "Lists"),
        ("Ingredients", "Storage", {"Pantry", "Cupboard", "Fridge",
                                    "Freezer"}, "storage"),
        ("Recipes", "Category",
         set(values("Lists", "Recipe categories")), "Lists"),
        ("Recipes", "Cuisine", set(values("Lists", "Cuisines")), "Lists"),
    ]
    for tab, header, allowed, source in checks:
        used = set(values(tab, header))
        missing = sorted(used - allowed)
        if missing:
            problems.append(f"{tab} / {header}: {len(missing)} value(s) not "
                            f"found on {source}: {missing[:6]}")

    # ---- meal plan sanity ------------------------------------------------------
    plan_dates = values("Meal Plan", "Date")
    bad_dates = [d for d in plan_dates if not hasattr(d, "year")]
    if bad_dates:
        problems.append(f"Meal Plan: {len(bad_dates)} date cells are not "
                        f"real dates")
    nutrition_dates = values("Nutrition", "Date")
    bad_nutrition = [d for d in nutrition_dates if not hasattr(d, "year")]
    if bad_nutrition:
        problems.append(f"Nutrition: {len(bad_nutrition)} date cells are not "
                        f"real dates")

    # ---- recipe ingredients -----------------------------------------------------
    ri = sheet("Recipe Ingredients")
    if ri is not None:
        recipe_col = col(ri, "Recipe")
        qty_col = col(ri, "Quantity")
        recipe_names = set()
        per_recipe = defaultdict(int)
        for row in range(5, 5 + args.rows):
            name = ri.cell(row=row, column=recipe_col).value \
                if recipe_col else None
            if not name:
                continue
            recipe_names.add(name)
            qty = ri.cell(row=row, column=qty_col).value if qty_col else None
            if isinstance(qty, (int, float)) and qty > 0:
                per_recipe[name] += 1
        without_ingredients = sorted(recipes - recipe_names)
        if without_ingredients:
            problems.append(f"{len(without_ingredients)} recipes have no "
                            f"ingredients: {without_ingredients[:6]}")
        notes.append(f"{len(recipe_names)} recipes carry "
                     f"{sum(per_recipe.values())} ingredient lines")

    # ---- budget months line up with the spending tab ------------------------------
    budget_months = set(values("Budget", "Month"))
    spending_months = set(values("Grocery Spending", "Month"))
    orphan = sorted(budget_months - spending_months)
    if orphan:
        problems.append(f"Budget months with no spending recorded: {orphan}")

    # ---- report ---------------------------------------------------------------------
    for line in notes:
        print(f"  {line}")
    print()
    if problems:
        print(f"FAIL - {len(problems)} problem(s)")
        for problem in problems:
            print(f"  - {problem}")
        return 1
    print("PASS - the example data hangs together")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
