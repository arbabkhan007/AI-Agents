"""Cross-check the example records in a generated workbook.

The demo data has to join up, otherwise the dashboard opens full of zeros.
This reads the literal values out of the finished file and makes sure every
client, pet, service, staff member, booking and invoice reference points at
a record that really exists.

    python tools/check_sample_data.py output/test.xlsx
"""

from __future__ import annotations

import sys

import openpyxl

HEADER_ROW = 4
INCIDENT_HEADER_ROW = 17
TASK_HEADER_ROW = 38


def column_values(ws, header: str, header_row: int = HEADER_ROW) -> list:
    """Every non-empty literal value under the given header."""
    headers = {}
    for cell in ws[header_row]:
        if isinstance(cell.value, str):
            headers[cell.value.strip()] = cell.column
    if header not in headers:
        raise KeyError(f"{ws.title}: no column called {header!r} "
                       f"(have {sorted(headers)[:8]} ...)")
    col = headers[header]
    out = []
    for row in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(row=row, column=col).value
        if value is None:
            continue
        if isinstance(value, str) and (value.startswith("=") or not value):
            continue
        out.append(value)
    return out


def column_rows(ws, header: str, header_row: int = HEADER_ROW) -> dict:
    """`{row number: value}` for every non-empty literal under a header.

    Same as `column_values` but it keeps the row, which matters for
    columns that are mostly blank.
    """
    headers = {}
    for cell in ws[header_row]:
        if isinstance(cell.value, str):
            headers[cell.value.strip()] = cell.column
    if header not in headers:
        raise KeyError(f"{ws.title}: no column called {header!r}")
    col = headers[header]
    out = {}
    for row in range(header_row + 1, ws.max_row + 1):
        value = ws.cell(row=row, column=col).value
        if value is None or (isinstance(value, str)
                             and (not value or value.startswith("="))):
            continue
        out[row] = value
    return out


def main(path: str) -> int:
    wb = openpyxl.load_workbook(path)
    problems = []

    clients = set(column_values(wb["Clients"], "Client name"))
    services = set(column_values(wb["Services"], "Service name"))
    staff = set(column_values(wb["Staff"], "Name"))
    pet_names = column_values(wb["Pets"], "Pet name")
    pet_owners = column_values(wb["Pets"], "Owner (client)")
    pets = {f"{n} ({o})" for n, o in zip(pet_names, pet_owners)}

    # IDs are formulas, so rebuild them the same way the sheet does.
    booking_rows = len(column_values(wb["Bookings"], "Client"))
    bookings = {f"BK-{i:04d}" for i in range(1, booking_rows + 1)}
    invoice_rows = len(column_values(wb["Invoices"], "Client"))
    invoices = {f"INV-{1001 + i}" for i in range(invoice_rows)}

    def check(sheet, header, allowed, label, header_row=HEADER_ROW,
              optional=True):
        values = column_values(wb[sheet], header, header_row)
        if not optional and not values:
            problems.append(f"{sheet}.{header}: no example data")
        for value in values:
            if value not in allowed:
                problems.append(f"{sheet}.{header}: {value!r} is not a "
                                f"known {label}")

    print(f"clients={len(clients)} pets={len(pets)} services={len(services)} "
          f"staff={len(staff)} bookings={len(bookings)} "
          f"invoices={len(invoices)}")

    check("Pets", "Owner (client)", clients, "client", optional=False)
    check("Bookings", "Client", clients, "client", optional=False)
    check("Bookings", "Pet", pets, "pet reference", optional=False)
    check("Bookings", "Service", services, "service", optional=False)
    check("Bookings", "Assigned to", staff, "staff member")
    check("Visit Log", "Booking ID", bookings, "booking", optional=False)
    check("Visit Log", "Carer", staff, "staff member")
    check("Health Log", "Pet", pets, "pet reference")
    check("Health Log", "Given by", staff, "staff member")
    check("Messages", "Client", clients, "client")
    check("Invoices", "Client", clients, "client", optional=False)
    check("Invoices", "Booking ref", bookings, "booking")
    check("Payments", "Client", clients, "client", optional=False)
    check("Payments", "Invoice no.", invoices, "invoice", optional=False)
    check("Mileage", "Client / booking", clients | bookings,
          "client or booking")
    check("Mileage", "Driver", staff, "staff member")
    check("Expenses", "Staff member", staff, "staff member")
    check("Expenses", "Related client / booking", clients | bookings,
          "client or booking")
    check("Incidents", "Client", clients, "client", INCIDENT_HEADER_ROW)
    check("Incidents", "Pet", pets, "pet reference", INCIDENT_HEADER_ROW)
    check("Reminders", "Owner", staff, "staff member", TASK_HEADER_ROW)

    # payments must be posted against their own client's invoice
    inv_client = {f"INV-{1001 + i}": value for i, value in
                  enumerate(column_values(wb["Invoices"], "Client"))}
    pay_clients = column_values(wb["Payments"], "Client")
    pay_invoices = column_values(wb["Payments"], "Invoice no.")
    for client, invoice in zip(pay_clients, pay_invoices):
        if inv_client.get(invoice) != client:
            problems.append(f"Payments: {invoice} belongs to "
                            f"{inv_client.get(invoice)!r} but the payment "
                            f"is logged against {client!r}")

    # visits must match the client and pet on their booking
    bk_client = {f"BK-{i + 1:04d}": value for i, value in
                 enumerate(column_values(wb["Bookings"], "Client"))}
    bk_pet = {f"BK-{i + 1:04d}": value for i, value in
              enumerate(column_values(wb["Bookings"], "Pet"))}
    vs_booking = column_values(wb["Visit Log"], "Booking ID")
    vs_client = column_values(wb["Visit Log"], "Client")
    vs_pet = column_values(wb["Visit Log"], "Pet")
    for booking, client, pet in zip(vs_booking, vs_client, vs_pet):
        if bk_client.get(booking) != client:
            problems.append(f"Visit Log: {booking} is for "
                            f"{bk_client.get(booking)!r}, visit says "
                            f"{client!r}")
        if bk_pet.get(booking) != pet:
            problems.append(f"Visit Log: {booking} is for "
                            f"{bk_pet.get(booking)!r}, visit says "
                            f"{pet!r}")

    # an invoice raised from a booking must be addressed to that booking's
    # client, otherwise the invoice total is pulled from somebody else's job
    inv_booking = column_rows(wb["Invoices"], "Booking ref")
    inv_clients = column_rows(wb["Invoices"], "Client")
    for row, booking in inv_booking.items():
        client = inv_clients.get(row)
        if bk_client.get(booking) != client:
            problems.append(f"Invoices: INV-{1001 + row - 5} is addressed "
                            f"to {client!r} but {booking} belongs to "
                            f"{bk_client.get(booking)!r}")

    if problems:
        print(f"! {len(problems)} problem(s)")
        for problem in problems[:40]:
            print(f"  - {problem}")
        return 1
    print("OK - every example record joins up")
    return 0


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else "output/test.xlsx"
    raise SystemExit(main(target))
