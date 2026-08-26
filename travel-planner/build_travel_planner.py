"""Build the Travel Planner v2 workbook (Novality store edition).

- 29 sheets, every sheet protected; only input (yellow) cells are editable.
- All formula / layout cells are locked with the password: premium
- Document author / creator: Novality store
"""
import os
import openpyxl
from openpyxl import Workbook

import helpers as H
import sheets_a, sheets_b, sheets_c, sheets_d

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                   "Travel_Planner_v2_Novality_Store.xlsx")


def main():
    wb = Workbook()
    wb.remove(wb.active)

    sheets_a.build_home(wb)          # Home + snapshot + navigation
    # --- Pre-trip planning ---
    sheets_a.build_destination(wb)
    sheets_a.build_budget(wb)
    sheets_a.build_savings(wb)
    sheets_a.build_research(wb)
    sheets_a.build_visa(wb)
    sheets_a.build_booking(wb)
    # --- Packing & preparation + itinerary ---
    sheets_b.build(wb)
    # --- Expenses + food & activities ---
    sheets_c.build(wb)
    # --- Extras & memories ---
    sheets_d.build(wb)

    # Document properties — author: Novality store
    p = wb.properties
    p.creator = "Novality store"
    p.lastModifiedBy = "Novality store"
    p.title = "Travel Planner v2"
    p.subject = "Travel planning spreadsheet"
    p.description = ("Travel Planner version 2 — 28 planning sheets covering budget, "
                     "packing, itinerary, expenses, food and memories. "
                     "Input cells editable; formula cells locked. (c) Novality store")
    p.keywords = "travel, planner, budget, itinerary, packing, expenses, Novality store"
    p.category = "Travel"

    wb.active = 0
    wb.save(OUT)
    print("Saved:", OUT)
    print("Sheets:", len(wb.sheetnames))
    for i, n in enumerate(wb.sheetnames, 1):
        print(f"  {i:2d}. {n}")


if __name__ == "__main__":
    main()
