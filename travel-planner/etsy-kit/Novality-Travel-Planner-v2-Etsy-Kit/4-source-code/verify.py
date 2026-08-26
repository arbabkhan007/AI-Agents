"""Verify Travel Planner v2: protection, locked formulas, refs, properties."""
import re
import openpyxl

P = "Travel_Planner_v2_Novality_Store.xlsx"
wb = openpyxl.load_workbook(P)
names = set(wb.sheetnames)
print("sheets:", len(wb.sheetnames))
assert len(wb.sheetnames) == 29

problems = []
unlocked_formula = []
formula_count = 0
unlocked_count = 0
for ws in wb.worksheets:
    # 1. protection
    if not ws.protection.sheet:
        problems.append(f"{ws.title}: protection OFF")
    if not ws.protection.password:
        problems.append(f"{ws.title}: no password hash")
    dvs = len(ws.data_validations.dataValidation)
    for row in ws.iter_rows():
        for c in row:
            if c.value is None:
                continue
            locked = c.protection.locked is not False
            if not locked:
                unlocked_count += 1
            if c.data_type == "f":
                formula_count += 1
                if not locked:
                    unlocked_formula.append(f"{ws.title}!{c.coordinate}: {c.value}")
                for m in re.findall(r"'([^']+)'!", str(c.value)):
                    if m not in names:
                        problems.append(f"{ws.title}!{c.coordinate}: bad ref -> {m}")

print("formula cells:", formula_count)
print("unlocked (input) cells:", unlocked_count)
print("unlocked formula cells:", len(unlocked_formula), unlocked_formula[:5])
print("problems:", problems if problems else "none")

# 2. spot checks
bt = wb["Budget Tracker"]
assert bt["B15"].value == "=SUM(B5:B14)"
assert bt["D5"].value.startswith('=IF(COUNT')
assert not bt["B15"].protection.locked is False
assert bt["B5"].protection.locked is False, "budget input should be unlocked"

home = wb["Home"]
assert "HYPERLINK" in home["D11"].value
assert "'Destination Overview'" in home["B5"].value

s = wb["Savings Goal"]
assert str(s["A9"].value).startswith('=IF(OR(B5=')

e = wb["Expense Log"]
assert "VLOOKUP" in e["G5"].value and "'Currency Converter'" in e["G5"].value
assert e["C158"].value.startswith("=SUMIF")

t = wb["Trip Cost Summary"]
assert t["B7"].value == "='Expense Log'!L5"
assert t["B11"].value == "='Souvenir List'!E25"

# password hash present on every sheet
for ws in wb.worksheets:
    assert ws.protection.password and ws.protection.sheet

p = wb.properties
print("creator:", p.creator, "| title:", p.title, "| lastModifiedBy:", p.lastModifiedBy)
assert p.creator == "Novality store"

# dropdowns sanity
for name, want in (("Packing Checklist", "A5:A64"), ("Expense Log", "F5:F154"),
                   ("Restaurant Wishlist", "D5:D34")):
    ws = wb[name]
    rngs = [str(d.sqref) for d in ws.data_validations.dataValidation]
    print(name, "validations:", rngs)
    assert any(want in r for r in rngs)

print("ALL CHECKS PASSED ✔")
