"""Sheets C — Expense Tracking + Food & Activities groups."""
import helpers as H
from openpyxl.utils import get_column_letter

EXPC = "E76F51"
FOOD = "7FB069"

CATEGORIES = ["Flights", "Accommodation", "Food & Dining", "Local Transport",
              "Activities & Tours", "Shopping", "Travel Insurance", "Visas & Fees",
              "Souvenirs", "Miscellaneous"]

RATES = [
    ("US Dollar", "USD", 1.0), ("Euro", "EUR", 1.08), ("British Pound", "GBP", 1.27),
    ("Japanese Yen", "JPY", 0.0067), ("Canadian Dollar", "CAD", 0.73),
    ("Australian Dollar", "AUD", 0.66), ("Swiss Franc", "CHF", 1.12),
    ("Chinese Yuan", "CNY", 0.14), ("Mexican Peso", "MXN", 0.055),
    ("Thai Baht", "THB", 0.028), ("Indian Rupee", "INR", 0.012),
    ("South Korean Won", "KRW", 0.0007),
]


def build_converter(wb):
    ws = H.new_sheet(wb, "Currency Converter", EXPC, "💱 CURRENCY CONVERTER",
                     "Rate table + instant converter — the Expense Log reads this table", 7)
    H.put(ws, "A3", "Home currency:", font=H.F_LABEL, halign="right")
    H.input_cell(ws, "B3", "USD", halign="center")
    ws.merge_cells("C3:D3")
    H.put(ws, "C3", "Used for all totals across the planner", font=H.F_NOTE,
          halign="left", indent=1)
    H.table(ws, 4, [
        {"h": "Currency", "w": 22, "k": "in", "a": "left"},
        {"h": "Code", "w": 9, "k": "in", "a": "center"},
        {"h": "1 Code → Home", "w": 14, "k": "in", "f": "0.0000"},
        {"h": "Rate Updated On", "w": 16, "k": "in", "f": H.DATEF, "a": "center"},
    ], 12, zebra=False)
    for i, (name, code, rate) in enumerate(RATES):
        r = 5 + i
        H.input_cell(ws, f"A{r}", name, halign="left", indent=1)
        H.input_cell(ws, f"B{r}", code, halign="center")
        H.input_cell(ws, f"C{r}", rate, fmt="0.0000")
    H.dv(ws, "B5:B16", ",".join(c for _, c, _ in RATES))

    ws.column_dimensions["F"].width = 18
    ws.column_dimensions["G"].width = 14
    ws.merge_cells("F4:G4")
    H.put(ws, "F4", "⚡ QUICK CONVERTER", font=H.F_HEAD, fillc=H.TEAL,
          halign="center", border=True)
    for r, label in ((5, "Amount"), (6, "Currency code"), (7, "= In home currency")):
        H.put(ws, f"F{r}", label, font=H.F_LABEL, fillc=H.CREAM, halign="left",
              indent=1, border=True)
    H.input_cell(ws, "G5", fmt=H.MONEY)
    H.input_cell(ws, "G6", halign="center")
    H.calc_cell(ws, "G7",
                '=IF(OR(G5="",G6=""),"",G5*IFERROR(VLOOKUP(G6,$B$5:$C$16,2,FALSE),1))',
                fmt=H.MONEY)
    H.put(ws, "A18", "Rates are approximate — update column C before you travel. "
                     "The Expense Log uses this table to auto-convert.",
          font=H.F_NOTE, halign="left", indent=1)
    H.finish(ws, landscape=False)


def build_expense(wb):
    ws = H.new_sheet(wb, "Expense Log", EXPC, "🧾 DAILY EXPENSE LOG",
                     "150 rows — enter amounts in any currency, totals convert automatically", 12)
    H.table(ws, 4, [
        {"h": "Date", "w": 13, "k": "in", "f": H.DATEF, "a": "center"},
        {"h": "Day", "w": 7, "k": "in", "f": H.INTF, "a": "center"},
        {"h": "Category", "w": 18, "k": "in", "a": "left"},
        {"h": "Item / Description", "w": 32, "k": "in", "a": "left"},
        {"h": "Amount", "w": 13, "k": "in", "f": H.MONEY},
        {"h": "Currency", "w": 10, "k": "in", "a": "center"},
        {"h": "Amount (Home)", "w": 14, "k": "calc", "f": H.MONEY},
        {"h": "Payment", "w": 13, "k": "in", "a": "center"},
        {"h": "Notes", "w": 26, "k": "in", "a": "left"},
    ], 150)
    for r in range(5, 155):
        H.calc_cell(ws, f"G{r}",
                    f'=IF(E{r}="","",IF(F{r}="",E{r},'
                    f'IFERROR(E{r}*VLOOKUP(F{r},\'Currency Converter\'!$B$5:$C$16,2,FALSE),"")))',
                    fmt=H.MONEY)
    H.dv(ws, "C5:C154", ",".join(CATEGORIES))
    H.dv(ws, "H5:H154", "Cash,Card,Mobile App,Travel Money Card,Other")
    H.dv_formula(ws, "F5:F154", "='Currency Converter'!$B$5:$B$16")
    H.summary_box(ws, 4, 11, "📊 TOTALS", [
        ("Total Spent (Home)", "=SUM(G5:G154)", H.MONEY),
        ("Entries Logged", "=COUNT(E5:E154)", H.INTF),
        ("Avg Spend / Day", "=IFERROR(ROUND(L5/'Destination Overview'!B14,2),\"—\")", H.MONEY),
    ], width=20)
    H.REFS["exp_total"] = "'Expense Log'!L5"

    H.section(ws, 156, 9, "📊 SPENDING BY CATEGORY")
    ws.merge_cells("A157:B157")
    H.put(ws, "A157", "Category", font=H.F_HEAD, fillc=H.TEAL, halign="center", border=True)
    ws["B157"].fill = H.fl(H.TEAL); ws["B157"].border = H.B_ALL
    H.put(ws, "C157", "Total Spent", font=H.F_HEAD, fillc=H.TEAL, halign="center", border=True)
    for i, cat in enumerate(CATEGORIES):
        r = 158 + i
        ws.merge_cells(f"A{r}:B{r}")
        H.put(ws, f"A{r}", cat, font=H.F_LABEL, fillc=H.CREAM, halign="left",
              indent=1, border=True)
        ws[f"B{r}"].fill = H.fl(H.CREAM); ws[f"B{r}"].border = H.B_ALL
        H.calc_cell(ws, f"C{r}", f'=SUMIF($C$5:$C$154,A{r},$G$5:$G$154)', fmt=H.MONEY)
    ws.merge_cells("A168:B168")
    H.put(ws, "A168", "TOTAL", font=H.F_TOTAL, fillc=H.NAVY, halign="left",
          indent=1, border=True)
    ws["B168"].fill = H.fl(H.NAVY); ws["B168"].border = H.B_ALL
    H.calc_cell(ws, "C168", "=SUM(C158:C167)", fmt=H.MONEY, font=H.F_TOTAL,
                fillc=H.NAVY)
    ws.auto_filter.ref = "A4:I154"
    H.finish(ws, freeze="A5", title_rows="4:4")


TIPS = [
    ("🇺🇸 USA", "15–25%", "10–15%", "$1–5 / bag", "10–20%", "Tipping is expected almost everywhere"),
    ("🇨🇦 Canada", "15–20%", "10–15%", "$1–5 / bag", "10–20%", "Similar to the US"),
    ("🇲🇽 Mexico", "10–15%", "Round up", "20–50 MXN", "10–15%", "Cash preferred, in pesos"),
    ("🇬🇧 United Kingdom", "10–12.5% if no service", "Round up", "£1–2 / bag", "10%", "Check bill — service often included"),
    ("🇫🇷 France", "Round up / 5%", "Round up", "€1–2 / bag", "10–15%", "Service compris — small tips welcome"),
    ("🇮🇹 Italy", "5–10% if no servizio", "Round up", "€1–2 / bag", "10%", "Coperto charge is not a tip"),
    ("🇪🇸 Spain", "Round up", "Round up", "€1–2 / bag", "€5–10", "Tipping is modest and casual"),
    ("🇩🇪 Germany", "5–10%, round up", "Round up", "€1–3 / bag", "€5–10", "Tell the server the total incl. tip"),
    ("🇯🇵 Japan", "No tip", "No tip", "No tip", "No tip", "Tipping can cause confusion — a bow works"),
    ("🇨🇳 China", "No tip", "No tip", "Sometimes expected", "Small gift", "Not customary in local spots"),
    ("🇹🇭 Thailand", "10% at nicer spots", "Round up", "20–50 THB", "100–300 THB", "Small tips appreciated"),
    ("🇦🇺 Australia", "10% optional", "Round up", "Optional", "10% optional", "Not required, always appreciated"),
]


def build_tipping(wb):
    ws = H.new_sheet(wb, "Tipping Guide", EXPC, "💁 TIPPING GUIDE",
                     "Country-by-country customs — never guess at the table again", 6)
    H.table(ws, 4, [
        {"h": "Country / Region", "w": 20, "k": "in", "a": "left"},
        {"h": "Restaurants", "w": 20, "k": "in", "a": "center"},
        {"h": "Taxis", "w": 14, "k": "in", "a": "center"},
        {"h": "Hotels", "w": 15, "k": "in", "a": "center"},
        {"h": "Tour Guides", "w": 14, "k": "in", "a": "center"},
        {"h": "Notes", "w": 38, "k": "in", "a": "left"},
    ], 16, zebra=False)
    for i, row in enumerate(TIPS):
        r = 5 + i
        for j, val in enumerate(row):
            H.input_cell(ws, f"{get_column_letter(1+j)}{r}", val,
                         halign="left" if j in (0, 5) else "center",
                         font=H.F_LABEL if j == 0 else H.F_BODY)
    H.put(ws, "A22", "Customs change fast — double-check for your exact destination. "
                     "Blank rows below are for your own research.",
          font=H.F_NOTE, halign="left", indent=1)
    H.finish(ws, freeze="A5")


def build_summary(wb):
    ws = H.new_sheet(wb, "Trip Cost Summary", EXPC, "📉 TRIP COST SUMMARY",
                     "Every total in the planner, pulled into one dashboard", 3)
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 17
    ws.column_dimensions["C"].width = 48
    R = H.REFS
    H.section(ws, 4, 3, "💰 COST BREAKDOWN")
    breakdown = [
        ("Estimated budget (Budget Tracker)", "='Budget Tracker'!B15", "What you plan to spend"),
        ("Actual — budget sheet (Budget Tracker)", "='Budget Tracker'!C15", "Per-category actuals you typed"),
        ("Logged expenses (Expense Log)", f"={R['exp_total']}", "Every line of the daily expense log"),
        ("Transportation booked (Transportation Log)", f"={R['trans_total']}", "Flights, trains, taxis…"),
        ("Accommodation booked (Accommodation)", f"={R['accom_total']}", "Nights × rate for all stays"),
        ("Itinerary activities (Daily Itinerary)", f"={R['itin_cost']}", "Costs planned in the itinerary"),
        ("Souvenirs & gifts actual (Souvenir List)", "='Souvenir List'!E25", "What you actually spent on gifts"),
    ]
    r = 5
    for label, f, note in breakdown:
        H.put(ws, f"A{r}", label, font=H.F_LABEL, fillc=H.CREAM, halign="left",
              indent=1, border=True)
        H.calc_cell(ws, f"B{r}", f, fmt=H.MONEY)
        H.put(ws, f"C{r}", note, font=H.F_NOTE, halign="left", indent=1, border=True)
        r += 1

    H.section(ws, 13, 3, "🩺 BUDGET HEALTH")
    health = [
        ("Total Estimated", "=B5", H.MONEY),
        ("Total Actual (budget sheet)", "=B6", H.MONEY),
        ("Remaining vs Budget", '=IF(B5=0,"",B5-B6)', H.MONEY),
        ("% of Budget Used", '=IF(B5=0,"",B6/B5)', H.PCT),
    ]
    r = 14
    for label, f, fmt in health:
        H.put(ws, f"A{r}", label, font=H.F_LABEL, fillc=H.CREAM, halign="left",
              indent=1, border=True)
        H.calc_cell(ws, f"B{r}", f, fmt=fmt)
        r += 1
    H.databar(ws, "B17")
    H.put(ws, "A18", "Status", font=H.F_LABEL, fillc=H.CREAM, halign="left",
          indent=1, border=True)
    H.calc_cell(ws, "B18",
                '=IF(B5=0,"Enter your budget to begin",IF(B6>B5,'
                '"⚠ Over budget by "&TEXT(B6-B5,"#,##0.00"),'
                '"✓ On track — "&TEXT(B5-B6,"#,##0.00")&" to spare"))',
                halign="center")
    ws.merge_cells("A20:C20")
    H.put(ws, "A20", "Log day-to-day spending in the Expense Log — totals flow into this page automatically.",
          font=H.F_NOTE, halign="left", indent=1)
    H.finish(ws, landscape=False)


def build_restaurant(wb):
    ws = H.new_sheet(wb, "Restaurant Wishlist", FOOD, "🍽️ RESTAURANT WISHLIST",
                     "Save every recommendation — book, visit, rate", 13)
    H.table(ws, 4, [
        {"h": "Restaurant", "w": 26, "k": "in", "a": "left"},
        {"h": "Cuisine", "w": 15, "k": "in", "a": "left"},
        {"h": "Address / Area", "w": 26, "k": "in", "a": "left"},
        {"h": "Price", "w": 9, "k": "in", "a": "center"},
        {"h": "Reservation", "w": 13, "k": "in", "a": "center"},
        {"h": "Res. Date", "w": 13, "k": "in", "f": H.DATEF, "a": "center"},
        {"h": "Priority", "w": 10, "k": "in", "a": "center"},
        {"h": "Visited?", "w": 10, "k": "in", "a": "center"},
        {"h": "Rating", "w": 9, "k": "in", "a": "center"},
        {"h": "Notes", "w": 26, "k": "in", "a": "left"},
    ], 30)
    H.dv(ws, "D5:D34", "$,$$,$$$,$$$$")
    H.dv(ws, "E5:E34", "Not Needed,To Book,Booked,Cancelled")
    H.dv(ws, "G5:G34", "High,Medium,Low")
    H.dv(ws, "H5:H34", "✓,✗")
    H.dv_num(ws, "I5:I34")
    H.summary_box(ws, 4, 12, "📊 SUMMARY", [
        ("Reservations Booked", '=COUNTIF(E5:E34,"Booked")', H.INTF),
        ("Places Visited", '=COUNTIF(H5:H34,"✓")', H.INTF),
        ("Avg Rating", '=IFERROR(ROUND(AVERAGE(I5:I34),1),"—")', H.DEC1),
    ], width=20)
    ws.auto_filter.ref = "A4:J34"
    H.finish(ws, freeze="A5")


def build_activity(wb):
    ws = H.new_sheet(wb, "Activity Planner", FOOD, "🎟️ ACTIVITY & EXCURSION PLANNER",
                     "From 💡 idea to ✅ booked — dates, costs and confirmations", 13)
    H.table(ws, 4, [
        {"h": "Activity / Excursion", "w": 32, "k": "in", "a": "left"},
        {"h": "Date", "w": 13, "k": "in", "f": H.DATEF, "a": "center"},
        {"h": "Time", "w": 9, "k": "in", "a": "center"},
        {"h": "Duration", "w": 11, "k": "in", "a": "center"},
        {"h": "Location / Meet Point", "w": 24, "k": "in", "a": "left"},
        {"h": "Cost", "w": 12, "k": "in", "f": H.MONEY},
        {"h": "Status", "w": 12, "k": "in", "a": "center"},
        {"h": "Confirmation #", "w": 16, "k": "in", "a": "center"},
        {"h": "Paid?", "w": 9, "k": "in", "a": "center"},
        {"h": "Notes", "w": 26, "k": "in", "a": "left"},
    ], 30)
    H.dv(ws, "G5:G34", "Idea,To Book,Booked,Cancelled")
    H.dv(ws, "I5:I34", "Yes,No,Deposit")
    H.summary_box(ws, 4, 12, "📊 SUMMARY", [
        ("Booked", '=COUNTIF(G5:G34,"Booked")', H.INTF),
        ("Total Cost", "=SUM(F5:F34)", H.MONEY),
        ("Still Just Ideas", '=COUNTIF(G5:G34,"Idea")', H.INTF),
    ], width=18)
    ws.auto_filter.ref = "A4:J34"
    H.finish(ws, freeze="A5")


PHRASES = [
    "Hello / Hi", "Good morning", "Good evening", "Please", "Thank you",
    "Yes / No", "Excuse me / Sorry", "How much does it cost?",
    "Where is the bathroom?", "Can you help me?", "I need a doctor",
    "Do you speak English?", "I would like…", "The check, please",
    "Cheers! (toast)", "Goodbye",
]


def build_phrases(wb):
    ws = H.new_sheet(wb, "Local Phrases", FOOD, "🗣️ LOCAL PHRASES",
                     "16 essentials — fill them in before you land", 4)
    H.table(ws, 4, [
        {"h": "English", "w": 26, "k": "in", "a": "left"},
        {"h": "Local Phrase", "w": 28, "k": "in", "a": "left"},
        {"h": "Pronunciation", "w": 26, "k": "in", "a": "left"},
        {"h": "Notes / When to Use", "w": 36, "k": "in", "a": "left"},
    ], 22)
    for i, p in enumerate(PHRASES):
        H.input_cell(ws, f"A{5+i}", p, font=H.F_LABEL, halign="left", indent=1)
    H.put(ws, "A28", "Hotel front desks love helping with pronunciation — ask!",
          font=H.F_NOTE, halign="left", indent=1)
    H.finish(ws, freeze="A5")


ETIQUETTE_TOPICS = [
    "Greetings & introductions", "Dining & table manners", "Tipping customs",
    "Dress code (temples, churches, cities)", "Religious & sacred sites",
    "Photography rules", "Public transport etiquette", "Gift-giving",
    "Punctuality & time", "Gestures & body language",
]


def build_etiquette(wb):
    ws = H.new_sheet(wb, "Etiquette Notes", FOOD, "🤝 CULTURAL ETIQUETTE NOTES",
                     "Blend in, show respect — research the do's and don'ts", 4)
    H.table(ws, 4, [
        {"h": "Topic", "w": 28, "k": "in", "a": "left"},
        {"h": "Customs & Tips", "w": 40, "k": "in", "a": "left"},
        {"h": "Do ✅", "w": 32, "k": "in", "a": "left"},
        {"h": "Don't 🚫", "w": 32, "k": "in", "a": "left"},
    ], 20)
    for i, t in enumerate(ETIQUETTE_TOPICS):
        H.input_cell(ws, f"A{5+i}", t, font=H.F_LABEL, halign="left", indent=1)
    H.finish(ws, freeze="A5")


def build(wb):
    build_converter(wb)
    build_expense(wb)
    build_tipping(wb)
    build_summary(wb)
    build_restaurant(wb)
    build_activity(wb)
    build_phrases(wb)
    build_etiquette(wb)
