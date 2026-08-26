"""Shared styling + helper utilities for the Travel Planner v2 workbook."""
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import DataBarRule
from openpyxl.worksheet.properties import PageSetupProperties

# ---------------------------------------------------------------- palette ---
NAVY   = "1D3557"
TEAL   = "2A9D8F"
SAND   = "E9C46A"
CORAL  = "E76F51"
WHITE  = "FFFFFF"
CREAM  = "F3F8F6"
INPUT  = "FFF8E1"   # yellow-ish = "type here"
CALC   = "E3ECF5"   # blue-gray  = "automatic / locked"
LINE   = "BFCFD6"
GRAYTX = "5F7470"

MONEY = "#,##0.00"
DATEF = "dd mmm yyyy"
PCT   = "0%"
INTF  = "#,##0"
DEC1  = "0.0"

UNLOCKED = Protection(locked=False)
LOCKED = Protection(locked=True)
_thin = Side(style="thin", color=LINE)
B_ALL = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)

F_TITLE = Font(bold=True, size=16, color=WHITE)
F_SUB   = Font(italic=True, size=10, color=WHITE)
F_HEAD  = Font(bold=True, size=10, color=WHITE)
F_BODY  = Font(size=10)
F_LABEL = Font(bold=True, size=10, color=NAVY)
F_CALC  = Font(bold=True, size=10, color=NAVY)
F_NOTE  = Font(italic=True, size=9, color=GRAYTX)
F_LINK  = Font(bold=True, size=10, color="1F7A70", underline="single")
F_TOTAL = Font(bold=True, size=10, color=WHITE)

REFS = {}   # cross-sheet cell references collected while building


def fl(h):
    return PatternFill("solid", fgColor=h)


def put(ws, coord, value=None, font=None, fillc=None, fmt=None, halign=None,
        valign="center", wrap=False, unlocked=False, border=False, indent=0):
    c = ws[coord]
    if value is not None:
        c.value = value
    c.font = font or F_BODY
    if fillc:
        c.fill = fl(fillc)
    if fmt:
        c.number_format = fmt
    c.alignment = Alignment(horizontal=halign, vertical=valign,
                            wrap_text=wrap, indent=indent)
    if border:
        c.border = B_ALL
    if unlocked:
        c.protection = UNLOCKED
    return c


def input_cell(ws, coord, value=None, fmt=None, halign=None, wrap=False, font=None,
               indent=0):
    return put(ws, coord, value, fillc=INPUT, fmt=fmt, halign=halign,
               wrap=wrap, unlocked=True, border=True, font=font, indent=indent)


def calc_cell(ws, coord, formula, fmt=None, halign="center", font=None, fillc=CALC):
    c = put(ws, coord, formula, font=font or F_CALC, fillc=fillc,
            fmt=fmt, halign=halign, border=True)
    c.protection = LOCKED          # formula cells are ALWAYS locked
    return c


def box(ws, r1, c1, r2, c2, fillc=None, unlocked=False):
    """Bordered (optionally merged) region."""
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = B_ALL
            if fillc:
                cell.fill = fl(fillc)
            if unlocked:
                cell.protection = UNLOCKED
    if (r1, c1) != (r2, c2):
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def banner(ws, ncols, title, subtitle):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    put(ws, "A1", title, font=F_TITLE, fillc=NAVY, halign="left", indent=1)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    put(ws, "A2", subtitle, font=F_SUB, fillc=TEAL, halign="left", indent=1)
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 6


def new_sheet(wb, name, tab, title, subtitle, ncols):
    ws = wb.create_sheet(title=name)
    ws.sheet_properties.tabColor = tab
    banner(ws, ncols, title, subtitle)
    return ws


def section(ws, row, ncols, text, start_col=1, color=SAND):
    c1 = start_col
    c2 = start_col + ncols - 1
    ws.merge_cells(start_row=row, start_column=c1, end_row=row, end_column=c2)
    L1 = get_column_letter(c1)
    put(ws, f"{L1}{row}", text, font=F_LABEL, fillc=color, halign="left",
        indent=1, border=True)
    ws.row_dimensions[row].height = 20


def table(ws, top, specs, nrows, start_col=1, zebra=True):
    """specs: list of dicts {h: header, w: width, k: 'in'|'calc'|'lock',
    f: number format, a: halign}.  Returns the last body row."""
    for j, sp in enumerate(specs):
        col = get_column_letter(start_col + j)
        ws.column_dimensions[col].width = sp["w"]
        put(ws, f"{col}{top}", sp["h"], font=F_HEAD, fillc=TEAL,
            halign="center", wrap=True, border=True)
    ws.row_dimensions[top].height = 22
    for r in range(top + 1, top + 1 + nrows):
        for j, sp in enumerate(specs):
            col = get_column_letter(start_col + j)
            if sp["k"] == "in":
                put(ws, f"{col}{r}", fillc=INPUT, fmt=sp.get("f"),
                    halign=sp.get("a"), unlocked=True, border=True)
            elif sp["k"] == "calc":
                put(ws, f"{col}{r}", font=F_CALC, fillc=CALC, fmt=sp.get("f"),
                    halign=sp.get("a", "center"), border=True)
            else:
                z = CREAM if (zebra and (r - top) % 2 == 0) else None
                put(ws, f"{col}{r}", fillc=z, fmt=sp.get("f"),
                    halign=sp.get("a"), border=True)
    return top + nrows


def summary_box(ws, top, col, title, items, width=22, vwidth=15, set_width=True):
    """items: list of (label, value_or_formula, fmt).  Returns last row used."""
    L1, L2 = get_column_letter(col), get_column_letter(col + 1)
    if set_width:
        ws.column_dimensions[L1].width = width
        ws.column_dimensions[L2].width = vwidth
    ws.merge_cells(f"{L1}{top}:{L2}{top}")
    put(ws, f"{L1}{top}", title, font=F_HEAD, fillc=TEAL, halign="center",
        border=True)
    ws.row_dimensions[top].height = 20
    r = top
    for label, value, fmt in items:
        r += 1
        put(ws, f"{L1}{r}", label, font=F_LABEL, fillc=CREAM, halign="left",
            indent=1, border=True)
        calc_cell(ws, f"{L2}{r}", value, fmt=fmt, halign="center")
    return r


def dv(ws, rng, options):
    d = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
    ws.add_data_validation(d)
    d.add(rng)
    return d


def dv_formula(ws, rng, ref):
    d = DataValidation(type="list", formula1=ref, allow_blank=True)
    ws.add_data_validation(d)
    d.add(rng)
    return d


def dv_num(ws, rng, lo=1, hi=5):
    d = DataValidation(type="whole", operator="between",
                       formula1=str(lo), formula2=str(hi), allow_blank=True)
    ws.add_data_validation(d)
    d.add(rng)
    return d


def databar(ws, rng, color="2A9D8F"):
    ws.conditional_formatting.add(
        rng, DataBarRule(start_type="num", start_value=0,
                         end_type="num", end_value=1,
                         color=color, showValue=True))


def protect(ws, pwd="premium"):
    p = ws.protection
    p.password = pwd
    p.sheet = True
    p.selectLockedCells = False    # False = action ALLOWED
    p.selectUnlockedCells = False
    p.formatColumns = False
    p.formatRows = False
    p.autoFilter = False


def finish(ws, landscape=True, freeze=None, title_rows=None):
    ws.sheet_view.showGridLines = False
    if freeze:
        ws.freeze_panes = freeze
    ws.oddFooter.left.text = "Travel Planner v2 - Novality store"
    ws.oddFooter.left.size = 8
    ws.oddFooter.left.color = "888888"
    ws.oddFooter.right.text = "Page &P of &N"
    ws.oddFooter.right.size = 8
    ws.oddFooter.right.color = "888888"
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    if title_rows:
        ws.print_title_rows = title_rows
    protect(ws)
