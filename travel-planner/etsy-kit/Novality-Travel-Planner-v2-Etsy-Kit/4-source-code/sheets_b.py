"""Sheets B — Packing & Preparation + Itinerary & Daily Planning groups."""
import helpers as H
from openpyxl.utils import get_column_letter

PACKG = "457B9D"
ITIN = "C99700"

PACK_ITEMS = [
    ("Clothing", "T-shirts & tops", None), ("Clothing", "Pants / jeans", None),
    ("Clothing", "Shorts", None), ("Clothing", "Dresses / going-out outfits", None),
    ("Clothing", "Underwear & socks", "7"), ("Clothing", "Pajamas", None),
    ("Clothing", "Swimwear", None), ("Clothing", "Light jacket / raincoat", None),
    ("Clothing", "Comfortable walking shoes", None),
    ("Clothing", "Formal outfit (1 set)", None), ("Clothing", "Hat / cap", None),
    ("Clothing", "Scarf / gloves (if cold)", None),
    ("Toiletries", "Toothbrush & toothpaste", None),
    ("Toiletries", "Shampoo & conditioner", None),
    ("Toiletries", "Soap / body wash", None), ("Toiletries", "Deodorant", None),
    ("Toiletries", "Skincare (AM / PM)", None), ("Toiletries", "Sunscreen SPF 30+", None),
    ("Toiletries", "Hairbrush / comb", None), ("Toiletries", "Razor & shaving cream", None),
    ("Toiletries", "Makeup", None), ("Toiletries", "Nail clippers", None),
    ("Toiletries", "Feminine care", None), ("Toiletries", "Quick-dry towel", None),
    ("Electronics", "Phone + charger", None), ("Electronics", "Power bank", None),
    ("Electronics", "Headphones / earbuds", None),
    ("Electronics", "Camera + memory card", None),
    ("Electronics", "Universal travel adapter", None),
    ("Electronics", "E-reader / tablet", None),
    ("Electronics", "Laptop + charger (optional)", None),
    ("Electronics", "Voltage converter (if needed)", None),
    ("Documents", "Passport", None), ("Documents", "Visa / eTA printout", None),
    ("Documents", "Travel insurance card", None), ("Documents", "Flight tickets", None),
    ("Documents", "Hotel confirmations", None),
    ("Documents", "Driver's license / IDP", None),
    ("Documents", "Cash & cards", None), ("Documents", "Copies of all documents", None),
    ("Health & Meds", "Prescription medications", None),
    ("Health & Meds", "Pain relievers", None), ("Health & Meds", "Allergy medicine", None),
    ("Health & Meds", "Stomach / motion sickness tablets", None),
    ("Health & Meds", "Band-aids & mini first-aid", None),
    ("Health & Meds", "Hand sanitizer", None), ("Health & Meds", "Face masks", None),
    ("Health & Meds", "Insect repellent", None), ("Health & Meds", "Aloe / sunburn relief", None),
    ("Miscellaneous", "Day bag / backpack", None),
    ("Miscellaneous", "Reusable water bottle", None),
    ("Miscellaneous", "Compact umbrella", None), ("Miscellaneous", "Laundry bag", None),
    ("Miscellaneous", "Travel snacks", None), ("Miscellaneous", "Travel pillow", None),
    ("Miscellaneous", "Books / entertainment", None),
    ("Miscellaneous", "Ziplock bags", None),
    ("Miscellaneous", "Foldable tote for souvenirs", None),
    ("Miscellaneous", "Earplugs & eye mask", None),
]


def build_packing(wb):
    ws = H.new_sheet(wb, "Packing Checklist", PACKG, "🧳 PACKING CHECKLIST",
                     "Clothing → toiletries → electronics → docs → meds — tick as you pack", 8)
    H.table(ws, 4, [
        {"h": "Category", "w": 18, "k": "in", "a": "left"},
        {"h": "Item", "w": 30, "k": "in", "a": "left"},
        {"h": "Qty", "w": 8, "k": "in", "f": H.INTF, "a": "center"},
        {"h": "Packed?", "w": 11, "k": "in", "a": "center"},
        {"h": "Notes", "w": 34, "k": "in", "a": "left"},
    ], 60)
    for i, (cat, item, qty) in enumerate(PACK_ITEMS):
        r = 5 + i
        H.input_cell(ws, f"A{r}", cat, halign="left", indent=1)
        H.input_cell(ws, f"B{r}", item, halign="left", indent=1)
        if qty:
            H.input_cell(ws, f"C{r}", int(qty), fmt=H.INTF, halign="center")
    H.dv(ws, "A5:A64", "Clothing,Toiletries,Electronics,Documents,Health & Meds,Miscellaneous")
    H.dv(ws, "D5:D64", "✓,✗")
    H.summary_box(ws, 4, 7, "📦 PROGRESS", [
        ("Total Items", "=COUNTA(B5:B64)", H.INTF),
        ("Packed", '=COUNTIF(D5:D64,"✓")', H.INTF),
        ("Still To Pack", "=H5-H6", H.INTF),
        ("% Packed", '=IF(H5=0,"",H6/H5)', H.PCT),
    ], width=16, vwidth=11)
    H.databar(ws, "H8")
    ws.auto_filter.ref = "A4:E64"
    H.finish(ws, freeze="A5", title_rows="4:4")


def build_weather(wb):
    ws = H.new_sheet(wb, "Weather Tracker", PACKG, "🌤️ WEATHER TRACKER",
                     "Check the forecast daily and plan outfits before you pack", 9)
    H.table(ws, 4, [
        {"h": "Date", "w": 13, "k": "in", "f": H.DATEF, "a": "center"},
        {"h": "Location", "w": 18, "k": "in", "a": "left"},
        {"h": "High °C", "w": 10, "k": "in", "f": H.DEC1},
        {"h": "Low °C", "w": 10, "k": "in", "f": H.DEC1},
        {"h": "Conditions", "w": 16, "k": "in", "a": "center"},
        {"h": "Outfit / What to Wear", "w": 42, "k": "in", "a": "left"},
    ], 20)
    H.dv(ws, "E5:E24", "☀ Sunny,⛅ Partly Cloudy,☁ Cloudy,🌧 Rain,⛈ Storm,❄ Snow,💨 Windy,🌫 Fog")
    H.summary_box(ws, 4, 8, "🌡️ SNAPSHOT", [
        ("Avg High", '=IF(COUNT(C5:C24)=0,"—",ROUND(AVERAGE(C5:C24),1))', H.DEC1),
        ("Avg Low", '=IF(COUNT(D5:D24)=0,"—",ROUND(AVERAGE(D5:D24),1))', H.DEC1),
        ("Coldest Low", '=IF(COUNT(D5:D24)=0,"—",MIN(D5:D24))', H.DEC1),
        ("Rainy Days", '=COUNTIF(E5:E24,"*Rain*")', H.INTF),
    ], width=13, vwidth=10)
    H.finish(ws, freeze="A5")


CARRY = [
    "Passport & travel documents", "Prescription medications",
    "Phone, chargers & power bank", "Change of clothes",
    "Toiletries ≤ 100 ml", "Snacks & empty water bottle",
    "Headphones / entertainment", "Sweater for the flight",
]
CHECKED = [
    "Clothing sets for each day", "Comfortable shoes (2nd pair)",
    "Full-size toiletries", "Hair dryer / styler",
    "Swimwear & beach items", "Foldable bag for souvenirs",
]


def build_bags(wb):
    ws = H.new_sheet(wb, "Carry-On vs Checked", PACKG, "👜 CARRY-ON vs CHECKED",
                     "Keep valuables & essentials with you — everything else goes below", 9)
    for start_col, title, items in ((1, "🎒 CARRY-ON (with you)", CARRY),
                                    (6, "🧳 CHECKED BAG (in the hold)", CHECKED)):
        H.section(ws, 4, 4, title, start_col=start_col)
        H.table(ws, 5, [
            {"h": "Item", "w": 28, "k": "in", "a": "left"},
            {"h": "Qty", "w": 7, "k": "in", "f": H.INTF, "a": "center"},
            {"h": "Packed?", "w": 10, "k": "in", "a": "center"},
            {"h": "Notes", "w": 24, "k": "in", "a": "left"},
        ], 20, start_col=start_col)
        cL = get_column_letter(start_col + 2)   # packed column
        iL = get_column_letter(start_col)
        for i, item in enumerate(items):
            H.input_cell(ws, f"{iL}{6+i}", item, halign="left", indent=1)
        H.dv(ws, f"{cL}6:{cL}25", "✓,✗")
        H.box(ws, 27, start_col, 27, start_col + 3, fillc=H.CALC)
        H.calc_cell(ws, f"{iL}27",
                    f'="Packed: "&COUNTIF({cL}6:{cL}25,"✓")&" of "&COUNTA({iL}6:{iL}25)',
                    halign="center")
    ws.column_dimensions["E"].width = 3
    H.finish(ws, freeze="A6")


TODO_ITEMS = [
    ("Notify bank of travel dates", "Money & Bank"), ("Buy / check travel insurance", "Money & Bank"),
    ("Get local currency or travel card", "Money & Bank"), ("Pay upcoming bills", "Money & Bank"),
    ("Arrange pet sitter / boarding", "Pets"), ("Arrange house sitter / neighbor check", "Home"),
    ("Hold mail or arrange pickup", "Home"), ("Water plants / plant sitter", "Plants"),
    ("Pause deliveries & subscriptions", "Home"), ("Empty fridge of perishables", "Home"),
    ("Take out trash / recycling", "Home"), ("Laundry — empty hamper", "Home"),
    ("Lock windows & doors, hide valuables", "Home"), ("Unplug appliances & electronics", "Home"),
    ("Set out-of-office / finish handover", "Work"), ("Charge devices & power bank", "Tech"),
    ("Download offline maps & entertainment", "Tech"), ("Print copies of key documents", "Travel Docs"),
    ("Check passport validity & visa", "Travel Docs"), ("Buy plug adapters / converters", "Tech"),
]


def build_todo(wb):
    ws = H.new_sheet(wb, "Pre-Trip To-Do", PACKG, "✅ TO-DO BEFORE YOU GO",
                     "The quiet list that saves the trip — bank, pets, plants, post", 8)
    H.table(ws, 4, [
        {"h": "Task", "w": 36, "k": "in", "a": "left"},
        {"h": "Category", "w": 15, "k": "in", "a": "left"},
        {"h": "Owner", "w": 14, "k": "in", "a": "left"},
        {"h": "Due Date", "w": 13, "k": "in", "f": H.DATEF, "a": "center"},
        {"h": "Done?", "w": 9, "k": "in", "a": "center"},
        {"h": "Notes", "w": 32, "k": "in", "a": "left"},
    ], 30)
    for i, (task, cat) in enumerate(TODO_ITEMS):
        r = 5 + i
        H.input_cell(ws, f"A{r}", task, halign="left", indent=1)
        H.input_cell(ws, f"B{r}", cat, halign="left", indent=1)
    H.dv(ws, "B5:B34", "Money & Bank,Home,Pets,Plants,Work,Tech,Health,Travel Docs,Other")
    H.dv(ws, "E5:E34", "✓,✗")
    H.summary_box(ws, 4, 8, "🏁 READINESS", [
        ("Done", '=COUNTIF(E5:E34,"✓")', H.INTF),
        ("Total Tasks", "=COUNTA(A5:A34)", H.INTF),
        ("% Ready", '=IF(H6=0,"",H5/H6)', H.PCT),
    ], width=13, vwidth=11)
    H.databar(ws, "H7")
    ws.auto_filter.ref = "A4:F34"
    H.finish(ws, freeze="A5")


def build_itinerary(wb):
    ws = H.new_sheet(wb, "Daily Itinerary", ITIN, "🗺️ DAY-BY-DAY ITINERARY",
                     "80 slots — time, activity, address, transport and cost per line", 12)
    H.table(ws, 4, [
        {"h": "Day", "w": 6, "k": "in", "a": "center"},
        {"h": "Date", "w": 13, "k": "in", "f": H.DATEF, "a": "center"},
        {"h": "Time / Block", "w": 13, "k": "in", "a": "center"},
        {"h": "Activity", "w": 34, "k": "in", "a": "left"},
        {"h": "Location / Address", "w": 28, "k": "in", "a": "left"},
        {"h": "Transport", "w": 14, "k": "in", "a": "center"},
        {"h": "Cost", "w": 12, "k": "in", "f": H.MONEY},
        {"h": "Booked?", "w": 10, "k": "in", "a": "center"},
        {"h": "Notes", "w": 28, "k": "in", "a": "left"},
    ], 80)
    r = 5
    for day in range(1, 6):
        for block in ("Morning", "Afternoon", "Evening"):
            H.input_cell(ws, f"A{r}", day, fmt=H.INTF, halign="center")
            H.input_cell(ws, f"C{r}", block, halign="center")
            r += 1
    H.dv(ws, "F5:F84", "Flight,Train,Bus / Coach,Metro / Subway,Taxi,Rideshare,Ferry,"
                       "Rental Car,Shuttle,Bike,Walk")
    H.dv(ws, "H5:H84", "Yes,No")
    H.summary_box(ws, 4, 11, "📊 SUMMARY", [
        ("Planned Activity Cost", "=SUM(G5:G84)", H.MONEY),
        ("Days Planned", '=IF(COUNT(A5:A84)=0,"—",MAX(A5:A84))', H.INTF),
    ], width=20)
    H.REFS["itin_cost"] = "'Daily Itinerary'!L5"
    ws.auto_filter.ref = "A4:I84"
    H.finish(ws, freeze="A5", title_rows="4:4")


def build_hourly(wb):
    ws = H.new_sheet(wb, "Hourly Schedule", ITIN, "⏰ HOURLY SCHEDULE GRID",
                     "One week, hour by hour — morning, afternoon and evening blocks", 9)
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 9
    for c in range(3, 10):
        ws.column_dimensions[get_column_letter(c)].width = 24
    heads = ["Block", "Time"] + [f"Day {i}" for i in range(1, 8)]
    for j, htxt in enumerate(heads):
        L = get_column_letter(1 + j)
        H.put(ws, f"{L}4", htxt, font=H.F_HEAD, fillc=H.TEAL, halign="center",
              border=True)
    ws.row_dimensions[4].height = 20
    H.put(ws, "B5", "Date →", font=H.F_NOTE, fillc=H.CREAM, halign="right", border=True)
    H.put(ws, "A5", "", fillc=H.CREAM, border=True)
    for c in range(3, 10):
        H.input_cell(ws, f"{get_column_letter(c)}5", fmt=H.DATEF, halign="center")

    blocks = [("☀ MORNING", 6, 11, "E7F6F2"), ("🌤 AFTERNOON", 12, 17, "FDF3D8"),
              ("🌙 EVENING", 18, 23, "FBE9E1")]
    for name, r1, r2, color in blocks:
        H.box(ws, r1, 1, r2, 1, fillc=color)
        ws.cell(row=r1, column=1).value = name
        ws.cell(row=r1, column=1).font = H.F_LABEL
        ws.cell(row=r1, column=1).alignment = H.Alignment(
            horizontal="center", vertical="center", text_rotation=90)
    for hour in range(6, 24):
        r = hour
        ws.row_dimensions[r].height = 20
        H.put(ws, f"B{r}", f"{hour:02d}:00", font=H.F_NOTE, fillc=H.CREAM,
              halign="center", border=True)
        for c in range(3, 10):
            H.put(ws, f"{get_column_letter(c)}{r}", fillc=H.WHITE, border=True,
                  unlocked=True)
    ws.row_dimensions[5].height = 18
    H.put(ws, "A25", "Tip: type plans into any hour. Morning ☀ 06–11 • Afternoon 🌤 12–17 • Evening 🌙 18–23.",
          font=H.F_NOTE, halign="left", indent=1)
    H.finish(ws, freeze="C6")


def build_route(wb):
    ws = H.new_sheet(wb, "Route Planner", ITIN, "🧭 MAP & ROUTE PLANNER",
                     "Order your stops, then sketch the route on the grid", 21)
    H.table(ws, 4, [
        {"h": "Stop #", "w": 8, "k": "lock", "a": "center"},
        {"h": "Place", "w": 26, "k": "in", "a": "left"},
        {"h": "Nearest Transit / Parking", "w": 26, "k": "in", "a": "left"},
        {"h": "Est. Time", "w": 12, "k": "in", "a": "center"},
        {"h": "Notes", "w": 30, "k": "in", "a": "left"},
    ], 10)
    for i in range(10):
        H.put(ws, f"A{5+i}", i + 1, font=H.F_NOTE, halign="center", border=True)

    for c in range(7, 22):
        ws.column_dimensions[get_column_letter(c)].width = 4.6
    H.put(ws, "U17", "N ⬆", font=H.F_LABEL, halign="center")
    for r in range(18, 33):
        ws.row_dimensions[r].height = 20
        for c in range(7, 22):
            H.put(ws, f"{get_column_letter(c)}{r}", fillc=H.WHITE, border=True,
                  unlocked=True)
    H.put(ws, "A34", "Sketch your route: mark ⬤ start, letter the stops A–J, shade 🟨 must-sees, "
                     "draw ➜ between cells.", font=H.F_NOTE, halign="left", indent=1)
    H.finish(ws, freeze="A5")


def build_transport(wb):
    ws = H.new_sheet(wb, "Transportation Log", ITIN, "🚉 TRANSPORTATION LOG",
                     "Every flight, train, bus and rideshare — times & confirmations", 14)
    H.table(ws, 4, [
        {"h": "Mode", "w": 14, "k": "in", "a": "center"},
        {"h": "Operator / Airline", "w": 20, "k": "in", "a": "left"},
        {"h": "From", "w": 16, "k": "in", "a": "left"},
        {"h": "To", "w": 16, "k": "in", "a": "left"},
        {"h": "Date", "w": 13, "k": "in", "f": H.DATEF, "a": "center"},
        {"h": "Depart", "w": 9, "k": "in", "a": "center"},
        {"h": "Arrive", "w": 9, "k": "in", "a": "center"},
        {"h": "Confirmation #", "w": 16, "k": "in", "a": "center"},
        {"h": "Seat / Terminal", "w": 13, "k": "in", "a": "center"},
        {"h": "Cost", "w": 12, "k": "in", "f": H.MONEY},
        {"h": "Notes", "w": 24, "k": "in", "a": "left"},
    ], 30)
    H.dv(ws, "A5:A34", "Flight,Train,Bus / Coach,Metro / Subway,Taxi,Rideshare,Ferry,"
                       "Rental Car,Shuttle,Bike,Walk")
    H.summary_box(ws, 4, 13, "📊 SUMMARY", [
        ("Total Transport Cost", "=SUM(J5:J34)", H.MONEY),
        ("Segments Logged", '=COUNTIF(A5:A34,"?*")', H.INTF),
    ], width=20)
    H.REFS["trans_total"] = "'Transportation Log'!N5"
    ws.auto_filter.ref = "A4:K34"
    H.finish(ws, freeze="A5")


FIELDS = [
    ("Name", None, None), ("Address", None, None), ("City", None, None),
    ("Phone", None, None), ("Confirmation #", None, None),
    ("Check-In Date", H.DATEF, None), ("Check-Out Date", H.DATEF, None),
    ("Nights", None, "calc"), ("Cost per Night", H.MONEY, None),
    ("Total Stay Cost", H.MONEY, "calc"), ("Wi-Fi Network", None, None),
    ("Wi-Fi Password", None, None), ("Check-In Time", None, None),
    ("Check-Out Time", None, None), ("Breakfast Included?", None, "dv"),
    ("Notes", None, "wide"),
]


def build_accommodation(wb):
    ws = H.new_sheet(wb, "Accommodation", ITIN, "🏨 ACCOMMODATION DETAILS",
                     "Up to 3 stays — dates, contacts, Wi-Fi and automatic totals", 6)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 2
    night_cells, total_cells, name_cells = [], [], []
    for k in range(3):
        head = 4 + k * 18
        H.section(ws, head, 3, f"🏨 STAY {k+1}")
        for i, (label, fmt, kind) in enumerate(FIELDS):
            r = head + 1 + i
            H.put(ws, f"A{r}", label, font=H.F_LABEL, fillc=H.CREAM,
                  halign="left", indent=1, border=True)
            H.box(ws, r, 2, r, 3, fillc=H.INPUT, unlocked=True)
            if fmt:
                ws[f"B{r}"].number_format = fmt
                ws[f"B{r}"].alignment = H.Alignment(horizontal="center", vertical="center")
            if label == "Name":
                name_cells.append(f"B{r}")
            if kind == "calc" and label == "Nights":
                H.box(ws, r, 2, r, 3, fillc=H.CALC)
                ci, co = head + 6, head + 7
                ws[f"B{r}"] = f'=IF(COUNT(B{ci}:B{co})<2,"",B{co}-B{ci})'
                ws[f"B{r}"].font = H.F_CALC
                ws[f"B{r}"].alignment = H.Alignment(horizontal="center", vertical="center")
                ws[f"B{r}"].protection = H.LOCKED
                night_cells.append(f"B{r}")
            if kind == "calc" and label == "Total Stay Cost":
                H.box(ws, r, 2, r, 3, fillc=H.CALC)
                n, cpn = head + 8, head + 9
                ws[f"B{r}"] = f'=IF(OR(B{n}="",B{cpn}=""),"",B{n}*B{cpn})'
                ws[f"B{r}"].font = H.F_CALC
                ws[f"B{r}"].number_format = H.MONEY
                ws[f"B{r}"].alignment = H.Alignment(horizontal="center", vertical="center")
                ws[f"B{r}"].protection = H.LOCKED
                total_cells.append(f"B{r}")
            if kind == "wide":
                ws.row_dimensions[r].height = 24
        # breakfast dropdown
        H.dv(ws, f"B{head+15}", "Yes,No")
    H.summary_box(ws, 4, 5, "📊 STAY SUMMARY", [
        ("Total Nights", "=" + "+".join(night_cells), H.INTF),
        ("Total Stay Cost", "=" + "+".join(total_cells), H.MONEY),
        ("Stays Planned", "=COUNTA(" + ",".join(name_cells) + ")", H.INTF),
    ], width=18, vwidth=14)
    H.REFS["accom_total"] = "'Accommodation'!F6"
    H.finish(ws)


def build(wb):
    build_packing(wb)
    build_weather(wb)
    build_bags(wb)
    build_todo(wb)
    build_itinerary(wb)
    build_hourly(wb)
    build_route(wb)
    build_transport(wb)
    build_accommodation(wb)
