"""
recalc_check.py - formula census for the workbook.

A full recalculation of a 60k+ formula workbook in Python is too slow to
run on every build (see the notes in the repo), so this tool does the next
best thing: it inventories every formula, lists the functions used, and
highlights the formula heavy sheets. Use find_circular_refs.py for cycle
detection and validate_social_workbook.py for the full structural pass.

Usage:
    python tools/recalc_check.py
    python tools/recalc_check.py "output/Social Media Content Planner.xlsx"
"""

from __future__ import annotations

import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DEFAULT = ROOT / "output" / "Social Media Content Planner.xlsx"

FUNCTION_RE = re.compile(r"([A-Z][A-Z0-9.]*)\s*\(")


def main(path: Path) -> int:
    if not path.exists():
        print(f"FAIL: workbook not found: {path}")
        return 1
    wb = load_workbook(path, data_only=False)
    total = 0
    per_sheet = Counter()
    functions = Counter()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if isinstance(value, str) and value.startswith("="):
                    total += 1
                    per_sheet[ws.title] += 1
                    stripped = re.sub(r'"[^"]*"', '""', value)
                    for fn in FUNCTION_RE.findall(stripped):
                        functions[fn] += 1

    print(f"{path.name}: {total:,} formulas across "
          f"{len(per_sheet)} sheets")
    print("\nHeaviest sheets:")
    for name, count in per_sheet.most_common(10):
        print(f"  {name:24s} {count:7,}")
    print("\nFunctions in use:")
    for fn, count in functions.most_common(30):
        print(f"  {fn:12s} {count:7,}")
    return 0


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT
    raise SystemExit(main(target))
