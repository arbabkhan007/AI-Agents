#!/usr/bin/env python3
"""Find circular references in a workbook.

Builds a dependency graph over formula cells only (static cells can't be part
of a cycle), expands referenced ranges, then reports every strongly connected
component (i.e. every group of formulas that feed each other in a loop).

Usage:
    /home/user/.venv/bin/python tools/find_circular_refs.py "output/Meal Planner & Grocery List.xlsx"
"""

import re
import sys
from collections import defaultdict

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

# ---------------------------------------------------------------------------
# Formula reference parsing
# ---------------------------------------------------------------------------

STRING_RE = re.compile(r'"(?:[^"]|"")*"')
# cross-sheet: 'Sheet Name'!A1 or Sheet!A1, optional range after the !
XSHEET_RE = re.compile(
    r"(?:'((?:[^']|'')+)'|([A-Za-z0-9_.][A-Za-z0-9_. ]*))!"
    r"(\$?[A-Z]{1,3}\$?\d+)(?::(\$?[A-Z]{1,3}\$?\d+))?"
)
# a cell ref that is not a function name or part of a longer identifier
CELL_RE = re.compile(r"(?<![A-Za-z0-9_.$:])(\$?[A-Z]{1,3}\$?\d+)(?![A-Za-z0-9_(])")
RANGE_RE = re.compile(
    r"(?<![A-Za-z0-9_.$:])(\$?[A-Z]{1,3}\$?\d+):(\$?[A-Z]{1,3}\$?\d+)(?![A-Za-z0-9_(])"
)


def parse_ref(token):
    """'$B$12' -> (row, col) 1-based."""
    m = re.fullmatch(r"\$?([A-Z]{1,3})\$?(\d+)", token)
    if not m:
        return None
    return int(m.group(2)), column_index_from_string(m.group(1))


def expand_range(a, b):
    r1, c1 = a
    r2, c2 = b
    if r1 > r2:
        r1, r2 = r2, r1
    if c1 > c2:
        c1, c2 = c2, c1
    return r1, c1, r2, c2


def extract_refs(formula, sheet_name, all_sheet_names):
    """Return list of (sheet, r1, c1, r2, c2) referenced by the formula."""
    # blank out string literals so refs inside text are not matched
    text = STRING_RE.sub(lambda m: '""', formula)
    refs = []
    consumed = []  # spans already handled (cross-sheet refs)

    for m in XSHEET_RE.finditer(text):
        name = m.group(1) if m.group(1) is not None else m.group(2)
        # normalise doubled quotes
        if m.group(1) is not None:
            name = name.replace("''", "'")
        if name not in all_sheet_names:
            continue
        a = parse_ref(m.group(3))
        b = parse_ref(m.group(4)) if m.group(4) else None
        if a is None:
            continue
        if b is not None:
            r1, c1, r2, c2 = expand_range(a, b)
        else:
            r1, c1 = a
            r2, c2 = r1, c1
        refs.append((name, r1, c1, r2, c2))
        consumed.append((m.start(3) - len(m.group(0)) + m.start(3) - m.start(3),
                         m.end()))
        # simpler: mask the whole match
        consumed[-1] = (m.start(), m.end())

    masked = list(text)
    for s, e in consumed:
        for i in range(s, e):
            masked[i] = " "
    masked = "".join(masked)

    # ranges first so single-cell regex does not double count corners
    spans = []
    for m in RANGE_RE.finditer(masked):
        a = parse_ref(m.group(1))
        b = parse_ref(m.group(2))
        if a is None or b is None:
            continue
        r1, c1, r2, c2 = expand_range(a, b)
        refs.append((sheet_name, r1, c1, r2, c2))
        spans.append((m.start(), m.end()))

    masked2 = list(masked)
    for s, e in spans:
        for i in range(s, e):
            masked2[i] = " "
    masked2 = "".join(masked2)

    for m in CELL_RE.finditer(masked2):
        a = parse_ref(m.group(1))
        if a is None:
            continue
        r1, c1 = a
        refs.append((sheet_name, r1, c1, r1, c1))

    return refs


# ---------------------------------------------------------------------------
# Graph + cycle detection
# ---------------------------------------------------------------------------

def find_cycles(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    all_sheet_names = set(wb.sheetnames)

    # node id per formula cell; formula text kept for reporting
    node_of = {}          # (sheet, row, col) -> node id
    node_info = []        # node id -> (sheet, row, col, formula)
    # per sheet, per column: sorted rows that contain formulas
    formula_rows = defaultdict(lambda: defaultdict(list))

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str) and v.startswith("="):
                    key = (ws.title, cell.row, cell.column)
                    node_of[key] = len(node_info)
                    node_info.append((ws.title, cell.row, cell.column, v))
                    formula_rows[ws.title][cell.column].append(cell.row)

    for sheet in formula_rows:
        for col in formula_rows[sheet]:
            formula_rows[sheet][col].sort()

    import bisect

    def formula_cells_in(sheet, r1, c1, r2, c2):
        cols = formula_rows.get(sheet, {})
        for col in range(c1, c2 + 1):
            rows = cols.get(col)
            if not rows:
                continue
            lo = bisect.bisect_left(rows, r1)
            hi = bisect.bisect_right(rows, r2)
            for i in range(lo, hi):
                yield (sheet, rows[i], col)

    edges = defaultdict(set)   # node -> set(nodes)
    n = len(node_info)
    for nid in range(n):
        sheet, row, col, formula = node_info[nid]
        for ref_sheet, r1, c1, r2, c2 in extract_refs(
                formula, sheet, all_sheet_names):
            # clamp huge ranges to a sane window is NOT done: exactness first
            for key in formula_cells_in(ref_sheet, r1, c1, r2, c2):
                dep = node_of[key]
                edges[nid].add(dep)

    # Tarjan SCC (iterative)
    index_counter = [0]
    index = {}
    lowlink = {}
    on_stack = set()
    stack = []
    sccs = []

    for start in range(n):
        if start in index:
            continue
        work = [(start, iter(edges.get(start, ())))]
        index[start] = lowlink[start] = index_counter[0]
        index_counter[0] += 1
        stack.append(start)
        on_stack.add(start)
        while work:
            node, it = work[-1]
            advanced = False
            for nxt in it:
                if nxt not in index:
                    index[nxt] = lowlink[nxt] = index_counter[0]
                    index_counter[0] += 1
                    stack.append(nxt)
                    on_stack.add(nxt)
                    work.append((nxt, iter(edges.get(nxt, ()))))
                    advanced = True
                    break
                elif nxt in on_stack:
                    lowlink[node] = min(lowlink[node], index[nxt])
            if advanced:
                continue
            work.pop()
            if work:
                parent = work[-1][0]
                lowlink[parent] = min(lowlink[parent], lowlink[node])
            if lowlink[node] == index[node]:
                comp = []
                while True:
                    w = stack.pop()
                    on_stack.discard(w)
                    comp.append(w)
                    if w == node:
                        break
                sccs.append(comp)

    cycles = []
    for comp in sccs:
        if len(comp) > 1:
            cycles.append(comp)
        elif comp[0] in edges.get(comp[0], ()):
            cycles.append(comp)  # self-reference

    print(f"formula cells: {n}")
    print(f"circular groups found: {len(cycles)}")
    for k, comp in enumerate(cycles, 1):
        print(f"\n=== cycle group {k}: {len(comp)} cell(s) ===")
        comp_sorted = sorted(comp, key=lambda i: (node_info[i][0],
                                                  node_info[i][1],
                                                  node_info[i][2]))
        for nid in comp_sorted[:40]:
            sheet, row, col, formula = node_info[nid]
            addr = f"{get_column_letter(col)}{row}"
            print(f"  [{sheet}] {addr}: {formula[:160]}")
        if len(comp_sorted) > 40:
            print(f"  ... and {len(comp_sorted) - 40} more cells")
    return cycles


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    find_cycles(sys.argv[1])
