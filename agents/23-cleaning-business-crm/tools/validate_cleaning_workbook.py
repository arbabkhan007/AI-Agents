#!/usr/bin/env python3
"""Validate a Cleaning Business CRM workbook.

Checks everything the previous products check, plus the lessons learned
the hard way:

* every formula has balanced parentheses and quotes (Excel silently
  DELETES unbalanced formulas and shows a repair log);
* no '=' leading the stored formula text (double-equals bug);
* no unreplaced {token} template leftovers;
* no modern functions (Excel 2010 safety);
* no #REF! anywhere;
* sheets protected with the password; input cells unlocked;
* document creator is Novality Store.

    python3 tools/validate_cleaning_workbook.py [file.xlsx] [--blank]
"""

from __future__ import annotations

import re
import sys
import zipfile
from pathlib import Path

import openpyxl

EXPECTED_SHEETS = [
    "Start Here", "Settings", "Dashboard", "Quote Calculator",
    "Clients", "Properties", "Leads", "Services", "Quotes", "Quote Lines",
    "Jobs", "Job Tasks", "Calendar", "Checklists", "Employees",
    "Staff Schedule", "Timesheets", "Job Costing", "Invoices",
    "Invoice Lines", "Payments", "Supplies", "Supply Usage", "Equipment",
    "Expenses", "Mileage", "Client Retention", "Complaints", "Incidents",
    "Safety", "Financial Plan", "Reports", "Print Quote", "Print Invoice",
    "Print Job Sheet", "Lists", "Help",
]

BANNED = [
    "XLOOKUP", "XMATCH", "IFS", "LET", "MINIFS", "MAXIFS",
    "TEXTJOIN", "CONCAT", "UNIQUE", "FILTER", "SORT", "SEQUENCE",
    "SWITCH", "LAMBDA", "BYROW", "BYCOL", "MAKEARRAY", "VSTACK",
    "HSTACK", "TEXTSPLIT", "TOROW", "TOCOL",
]

BANNED_RE = re.compile(
    r"(?<![A-Z0-9_.])(" + "|".join(BANNED) + r")\s*\(")

# cells that must remain editable formulas that the user overwrites
SEEDED_INPUTS = ["Calendar!C5", "Staff Schedule!C5", "Quote Calculator!C6"]

TOKEN_LEFTOVER = re.compile(r"\{[A-Za-z_][A-Za-z0-9_]*\}")


def paren_depth(formula: str):
    """Return (ok, reason) after scanning quotes and parens."""
    depth = 0
    in_str = False
    i = 0
    n = len(formula)
    while i < n:
        ch = formula[i]
        if in_str:
            if ch == '"':
                if i + 1 < n and formula[i + 1] == '"':
                    i += 2
                    continue
                in_str = False
        else:
            if ch == '"':
                in_str = True
            elif ch == "(":
                depth += 1
            elif ch == ")":
                depth -= 1
                if depth < 0:
                    return False, "more ')' than '('"
        i += 1
    if in_str:
        return False, "unclosed string quote"
    if depth != 0:
        return False, f"unbalanced parens (depth {depth} at end)"
    return True, ""


CELLTOKEN = r"(?:\$?[A-Z]{1,3}\$?\d+|[a-z_][a-z0-9_]*\{r\})"


def precedence_slip(formula: str):
    """Catch `A5-B5/C5` written where `(A5-B5)/C5` was meant.

    Only flags the risky shape where a *cell reference* minus another
    *cell reference* is immediately divided - numeric literals like the
    intentional `(...)*24-G6/60` are not flagged.
    """
    slim = re.sub(r'"[^"]*"', '""', formula)
    m = re.search(CELLTOKEN + r"\s*-\s*" + CELLTOKEN + r"\s*/", slim)
    if m:
        return True, f"possible precedence slip near '{m.group(0)}'"
    return False, ""


def raw_xml_problems(path: Path):
    """Open the xlsx as a zip and scan raw sheet XML.

    Catches what openpyxl's view hides: formula cells whose cached <v>
    is the literal string 'None' (written when value=None is passed to
    xlsxwriter). Excel reports those as "Repaired Records: Cell
    information" in its recovery log.
    """
    import zipfile
    out = []
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not (name.startswith("xl/worksheets/")
                    and name.endswith(".xml")):
                continue
            data = z.read(name).decode("utf-8", "replace")
            for bad in ("<v>None</v>", "<v>NaN</v>", "<v>Infinity</v>"):
                n = data.count(bad)
                if n:
                    out.append(f"{name}: {n} cached {bad} values")
            for m in re.finditer(r"<f[^>]*>([^<]*)</f>", data):
                if "{" in m.group(1) or "}" in m.group(1):
                    out.append(f"{name}: braces in raw formula "
                               f":: {m.group(1)[:70]}")
    return out


def validate(path: Path, blank: bool = False) -> int:
    problems = []
    problems += raw_xml_problems(path)
    wb = openpyxl.load_workbook(path, data_only=False)

    # ---------------------------------------------------------- metadata
    creator = (wb.properties.creator or "")
    modified_by = (wb.properties.lastModifiedBy or "")
    if creator != "Novality Store":
        problems.append(f"creator is '{creator}', expected 'Novality Store'")
    if modified_by and modified_by != "Novality Store":
        problems.append(f"lastModifiedBy is '{modified_by}'")

    # --------------------------------------------------------- sheet list
    if wb.sheetnames != EXPECTED_SHEETS:
        problems.append(f"sheet order mismatch: {wb.sheetnames}")
        missing = set(EXPECTED_SHEETS) - set(wb.sheetnames)
        extra = set(wb.sheetnames) - set(EXPECTED_SHEETS)
        if missing:
            problems.append(f"  missing: {sorted(missing)}")
        if extra:
            problems.append(f"  extra: {sorted(extra)}")

    # ----------------------------------------------------------- formulas
    total_formulas = 0
    total_cells = 0
    for ws in wb.worksheets:
        if not ws.protection.sheet:
            problems.append(f"{ws.title}: sheet is NOT protected")
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value is None:
                    continue
                total_cells += 1
                if not isinstance(value, str) or not value.startswith("="):
                    continue
                total_formulas += 1
                where = f"{ws.title}!{cell.coordinate}"

                if value.startswith("=="):
                    problems.append(f"{where}: formula starts with '=='")

                ok, reason = paren_depth(value)
                if not ok:
                    problems.append(f"{where}: {reason} :: {value[:110]}")

                slip, why = precedence_slip(value)
                if slip:
                    problems.append(f"{where}: {why}")

                if "#REF!" in value:
                    problems.append(f"{where}: contains #REF!")

                leftover = TOKEN_LEFTOVER.search(value)
                if leftover:
                    problems.append(
                        f"{where}: unreplaced template token "
                        f"'{leftover.group(0)}'")

                if "{" in value or "}" in value:
                    problems.append(
                        f"{where}: braces in formula (unrendered token) "
                        f":: {value[:80]}")

                upper = value.upper()
                for m in BANNED_RE.finditer(upper):
                    problems.append(f"{where}: banned function "
                                    f"'{m.group(1)}'")

    # ------------------------------------------------------- seeded inputs
    for ref in SEEDED_INPUTS:
        sheet, coord = ref.split("!")
        if sheet not in wb.sheetnames:
            problems.append(f"{ref}: sheet missing")
            continue
        ws = wb[sheet]
        cell = ws[coord]
        if cell.protection.locked:
            problems.append(f"{ref}: input cell is locked")
        if cell.value is None or (isinstance(cell.value, str)
                                  and not cell.value.startswith("=")):
            # a raw input default is fine too (e.g. quote date)
            pass

    # -------------------------------------------------------- defined names
    names = list(wb.defined_names)
    for needed in ["Clients_Names", "Employees_Names", "Jobs_Numbers",
                   "Invoices_Numbers", "Quotes_Numbers",
                   "Business_Name", "Hourly_Rate", "Minimum_Charge"]:
        if needed not in names:
            problems.append(f"defined name '{needed}' missing")

    # --------------------------------------------------------- validation
    n_validations = 0
    for ws in wb.worksheets:
        n_validations += len(ws.data_validations.dataValidation)
    if n_validations < 50:
        problems.append(f"only {n_validations} data validations found")

    # --------------------------------------------------------- zip sanity
    with zipfile.ZipFile(path) as zf:
        bad = zf.testzip()
        if bad:
            problems.append(f"corrupt member in xlsx zip: {bad}")

    # ------------------------------------------------------------- report
    label = "BLANK" if blank else "FILLED"
    print(f"[{label}] {path.name}")
    print(f"  sheets         : {len(wb.sheetnames)}")
    print(f"  formulas       : {total_formulas:,}")
    print(f"  populated cells: {total_cells:,}")
    print(f"  validations    : {n_validations}")
    print(f"  defined names  : {len(names)}")

    if problems:
        print(f"  PROBLEMS       : {len(problems)}")
        for p in problems[:40]:
            print(f"    - {p}")
        if len(problems) > 40:
            print(f"    ... and {len(problems) - 40} more")
        return 1
    print("  RESULT         : PASS")
    return 0


def main(argv=None) -> int:
    argv = list(sys.argv[1:] if argv is None else argv)
    forced_blank = "--blank" in argv
    argv = [a for a in argv if a != "--blank"]
    base = Path(__file__).resolve().parent.parent / "output"
    if argv:
        files = [Path(a) for a in argv]
        blank = True if forced_blank else None
    else:
        files = [
            base / "Cleaning Business CRM & Pricing Calculator.xlsx",
            base / "Cleaning Business CRM & Pricing Calculator (Blank).xlsx",
        ]
        blank = None
    rc = 0
    for i, f in enumerate(files):
        is_blank = blank if blank is not None else ("Blank" in f.name)
        rc |= validate(f, blank=is_blank)
    return rc


if __name__ == "__main__":
    raise SystemExit(main())
