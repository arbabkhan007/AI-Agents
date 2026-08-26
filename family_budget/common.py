"""Shared constants and helpers for building the Novality Family Budget workbook."""
import datetime as dt
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection


# ---- Theme ------------------------------------------------------------------
TITLE_FILL   = "0B3D2E"   # deep emerald
HEADER_FILL  = "1F7A5C"   # primary green
HEADER_FILL2 = "2E8B6E"
SECTION_FILL = "DCEFE4"   # soft green section header
INPUT_FILL   = "FFF9EC"   # warm cream input cells
CALC_FILL    = "EAF6EF"   # formula / auto calc cells
TOTAL_FILL   = "CBE6D4"
ALT_FILL     = "F4F8F5"
GOOD_FILL    = "C6EFCE"
WARN_FILL    = "FFEB9C"
BAD_FILL     = "FFC7CE"
GOOD_FONT    = "006100"
WARN_FONT    = "9C6500"
BAD_FONT     = "9C0006"

# ---- Styles -----------------------------------------------------------------
THIN = Side(style="thin", color="B7C9C0")
MED  = Side(style="medium", color="0B3D2E")
BOX  = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(bold=False, size=11, color="000000", italic=False):
    return Font(name="Calibri", size=size, bold=bold, color=color, italic=italic)


def fill(hexcolor):
    return PatternFill("solid", fgColor=hexcolor)


def align(horizontal="left", vertical="center", wrap=False):
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)


def cell(ws, row, col, value=None, bold=False, size=11, color="000000",
         bg=None, horizontal="left", vertical="center", wrap=False,
         numfmt=None, locked=True, border=BOX, italic=False):
    c = ws.cell(row=row, column=col)
    if value is not None:
        c.value = value
    c.font = font(bold=bold, size=size, color=color, italic=italic)
    if bg is not None:
        c.fill = fill(bg)
    c.alignment = align(horizontal=horizontal, vertical=vertical, wrap=wrap)
    if numfmt is not None:
        c.number_format = numfmt
    c.protection = Protection(locked=locked)
    if border is not None:
        c.border = border
    return c


# Currency / date / percent formats
USD  = '"$"#,##0.00'
USD0 = '"$"#,##0'
PCT  = '0.0%'
PCT0 = '0%'
DATE = 'mm/dd/yyyy'
NUM  = '#,##0'
DAYSTR = 'dddd'
PASSWORD = "premium"
# ---------------------------------------------------------------------------
# Tab (sheet) names
# ---------------------------------------------------------------------------
DASH = "🏠 Dashboard"
INC  = "💰 Income"
TXN  = "💳 Transactions"
BUD  = "📋 Budget"
BIL  = "🧾 Bills & Subs"
ACC  = "🏦 Accounts"
DBT  = "💸 Debt Manager"
GOL  = "🎯 Goals & Funds"
INV  = "📈 Savings & Invest"
HOM  = "🏡 Home & Mortgage"
VEH  = "🚗 Vehicles"
FAM  = "👨‍👩‍👧‍👦 Family & Kids"
HEA  = "🏥 Healthcare"
TAX  = "🇺🇸 Tax Center"
ANA  = "📊 Analytics"
CAL  = "📅 Calendar"
SET  = "⚙️ Settings"
INS  = "📖 Instructions"



def q(name):
    """Quote a sheet name for a formula reference."""
    return "'" + name + "'"


def T(sheet):
    """Return a quoted-sheet-name string usable in formulas, e.g. 'Transactions'."""
    return q(sheet)


# ---- Master budget structure --------------------------------------------------
BUDGET = {
    "Housing": ["Mortgage/Rent", "Property Tax", "HOA", "Home Insurance",
                "Repairs/Maintenance", "Utilities", "Other Housing"],
    "Transportation": ["Car Payments", "Gas", "Auto Insurance", "Maintenance/Repairs",
                       "Registration", "Parking/Tolls", "Public Transportation"],
    "Food": ["Groceries", "Restaurants", "Takeout/Delivery", "Coffee", "Other Food"],
    "Family": ["Childcare", "School", "Activities/Sports", "Clothing", "Allowances",
               "Baby Expenses", "College Savings", "Other Family"],
    "Healthcare": ["Health Insurance", "Dental", "Vision", "Prescriptions",
                   "Medical Expenses", "Other Healthcare"],
    "Lifestyle": ["Entertainment", "Subscriptions", "Travel", "Hobbies",
                  "Personal Care", "Shopping", "Other Lifestyle"],
    "Financial": ["Credit Cards", "Student Loans", "Auto Loans", "Personal Loans",
                  "Savings", "Investments", "Retirement", "Other Financial"],
    "Giving": ["Charity", "Gifts", "Tithe", "Other Giving"],
    "Pets": ["Food", "Vet", "Grooming", "Supplies", "Other Pets"],
    "Travel": ["Transport", "Lodging", "Food & Dining", "Activities", "Other Travel"],
    "Miscellaneous": ["Fees", "Gifts", "Other Misc"],
}

CATEGORIES = list(BUDGET.keys())
SUBCATEGORIES = [s for v in BUDGET.values() for s in v]

# Default monthly budget limits (USD) - edit these in the Budget tab.
DEFAULT_BUDGET = {
    ('Housing', 'Mortgage/Rent'): 1250,
    ('Housing', 'Property Tax'): 185,
    ('Housing', 'HOA'): 35,
    ('Housing', 'Home Insurance'): 90,
    ('Housing', 'Repairs/Maintenance'): 85,
    ('Housing', 'Utilities'): 235,
    ('Transportation', 'Car Payments'): 265,
    ('Transportation', 'Gas'): 140,
    ('Transportation', 'Auto Insurance'): 85,
    ('Transportation', 'Maintenance/Repairs'): 65,
    ('Transportation', 'Registration'): 25,
    ('Transportation', 'Parking/Tolls'): 35,
    ('Transportation', 'Public Transportation'): 20,
    ('Food', 'Groceries'): 435,
    ('Food', 'Restaurants'): 140,
    ('Food', 'Takeout/Delivery'): 45,
    ('Food', 'Coffee'): 35,
    ('Family', 'Childcare'): 585,
    ('Family', 'School'): 60,
    ('Family', 'Activities/Sports'): 75,
    ('Family', 'Clothing'): 65,
    ('Family', 'Allowances'): 25,
    ('Family', 'Baby Expenses'): 40,
    ('Family', 'College Savings'): 125,
    ('Healthcare', 'Health Insurance'): 290,
    ('Healthcare', 'Dental'): 20,
    ('Healthcare', 'Vision'): 15,
    ('Healthcare', 'Prescriptions'): 35,
    ('Healthcare', 'Medical Expenses'): 65,
    ('Lifestyle', 'Entertainment'): 90,
    ('Lifestyle', 'Subscriptions'): 60,
    ('Lifestyle', 'Travel'): 150,
    ('Lifestyle', 'Hobbies'): 45,
    ('Lifestyle', 'Personal Care'): 40,
    ('Lifestyle', 'Shopping'): 90,
    ('Financial', 'Credit Cards'): 150,
    ('Financial', 'Student Loans'): 165,
    ('Financial', 'Auto Loans'): 210,
    ('Financial', 'Personal Loans'): 75,
    ('Financial', 'Savings'): 210,
    ('Financial', 'Investments'): 125,
    ('Financial', 'Retirement'): 250,
    ('Giving', 'Charity'): 85,
    ('Giving', 'Gifts'): 35,
    ('Pets', 'Food'): 60,
    ('Pets', 'Vet'): 35,
    ('Pets', 'Supplies'): 20,
    ('Miscellaneous', 'Fees'): 25,
}

# ---- Settings dropdown lists --------------------------------------------------
FAMILY = ["Parent 1", "Parent 2", "Child 1", "Child 2", "Household", "Other"]
INCOME_SOURCES = ["Salary", "Hourly", "Bonus", "Overtime", "Commission",
                  "Freelance/Side Hustle", "Child Support", "Social Security",
                  "Retirement", "Government Benefits", "Other"]
PAY_FREQ = ["Weekly", "Biweekly", "Twice Monthly", "Monthly", "Irregular"]
ACCOUNT_TYPES = ["Checking", "Savings", "HYSA", "Cash", "Credit Card", "401(k)",
                 "403(b)", "IRA", "Roth IRA", "Brokerage", "529", "HSA", "FSA",
                 "CD", "Other Investment", "Other"]
PAY_METHODS = ["Checking", "Savings", "Credit Card", "Cash", "Autopay",
               "Debit Card", "Other"]
TX_TYPES = ["Expense", "Income", "Transfer"]
BILL_FREQ = ["Weekly", "Biweekly", "Twice Monthly", "Monthly", "Quarterly",
             "Annually", "One-time"]
YES_NO = ["Yes", "No"]
ESS_DISC = ["Essential", "Discretionary"]
MONTHS = ["January", "February", "March", "April", "May", "June", "July",
          "August", "September", "October", "November", "December"]
FILING = ["Single", "Married Filing Jointly", "Married Filing Separately",
          "Head of Household", "Qualifying Surviving Spouse"]
DEBT_TYPES = ["Credit Card", "Student Loan", "Auto Loan", "Personal Loan", "Mortgage", "Other"]
GOAL_TYPES = ["Emergency Fund", "Vacation", "Christmas", "Birthdays", "Car Repairs",
              "Home Repairs", "Insurance", "Property Taxes", "College", "Major Purchase", "Custom"]
INVEST_TYPES = ["401(k)", "403(b)", "IRA", "Roth IRA", "Brokerage", "HSA",
                "529", "CD", "HYSA", "Cash", "Emergency Fund"]


# ---------------------------------------------------------------------------
# Workbook-level helpers shared by every tab builder
# ---------------------------------------------------------------------------
LISTS = {}          # list display name -> Settings range string
from openpyxl.worksheet.protection import SheetProtection


def protect(sheet, allow_sort=True, allow_filter=True):
    sp = SheetProtection(sheet=True, password=PASSWORD, objects=True, scenarios=True,
                         formatCells=False, formatColumns=False, formatRows=False,
                         insertColumns=False, insertRows=False, insertHyperlinks=False,
                         deleteColumns=False, deleteRows=False, selectLockedCells=True,
                         selectUnlockedCells=True, sort=allow_sort, autoFilter=allow_filter,
                         pivotTables=False)
    sheet.protection = sp


def title_band(sheet, text, sub, span_cols):
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=span_cols)
    c = cell(sheet, 1, 1, text, bold=True, size=16, color="FFFFFF", bg=TITLE_FILL,
             horizontal="left", border=None, locked=False)
    c.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 30
    if sub:
        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=span_cols)
        c2 = cell(sheet, 2, 1, sub, size=9, italic=True, color="3D4A46",
                  bg="EAF3EE", border=None, locked=False)
        c2.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        sheet.row_dimensions[2].height = 28


def header_row(sheet, row, labels, bg=HEADER_FILL, start=1):
    for i, lab in enumerate(labels):
        cell(sheet, row, start + i, lab, bold=True, size=10, color="FFFFFF",
             bg=bg, horizontal="center", wrap=True, locked=False)


def add_dv(sheet, dv_type, formula1, ranges, allow_blank=True, error=True,
           prompt=None, error_style=None):
    # Normalize formula1: named ranges / cell ranges must not start with '='.
    if isinstance(formula1, str) and formula1.startswith("=") and not formula1.startswith('="'):
        formula1 = formula1[1:]
    dv = DataValidation(type=dv_type, formula1=formula1, allow_blank=allow_blank)
    if prompt:
        dv.promptTitle = prompt
        dv.prompt = prompt
    dv.showErrorMessage = error
    dv.showInputMessage = bool(prompt)
    dv.showDropDown = False
    if error_style:
        dv.errorStyle = error_style
    sheet.add_data_validation(dv)
    if isinstance(ranges, str):
        ranges = [x.strip() for x in ranges.split(",") if x.strip()]
    for r in ranges:
        dv.add(r)


def set_widths(sheet, widths):
    for i, w in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = w


def money_col(ws, col_letter, first, last, extra=USD):
    for r in range(first, last + 1):
        ws[f"{col_letter}{r}"].number_format = extra
