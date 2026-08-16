from pathlib import Path
from copy import copy
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "Interior_Design_Project_Hub.xlsx"
OUT.parent.mkdir(parents=True, exist_ok=True)

# Design system: calm navy + teal, warm input cells, and clear status colors.
NAVY = "1F2A44"
TEAL = "2F6F73"
GOLD = "D6A756"
INK = "263238"
WHITE = "FFFFFF"
PALE_TEAL = "EAF3F3"
PALE_BLUE = "EAF3F8"
INPUT = "FFF8E7"
LIGHT_GRAY = "F3F5F7"
MID_GRAY = "E7E6E6"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
RED = "FCE4D6"
GRAY = "E7E6E6"
BORDER = "D6DDE5"

thin_gray = Side(style="thin", color=BORDER)
medium_navy = Side(style="medium", color=NAVY)

money_fmt = '$#,##0.00;[Red]-$#,##0.00;-'
whole_money_fmt = '$#,##0;[Red]-$#,##0;-'
pct_fmt = '0.0%'
date_fmt = 'm/d/yy'

wb = Workbook()
ws = wb.active
ws.title = "Welcome"
wb.properties.title = "Interior Design Project Hub"
wb.properties.subject = "Editable interior design project management spreadsheet"
wb.properties.creator = "Interior Design Project Hub"
wb.properties.description = "Google Sheets and Microsoft Excel compatible interior design project tracker template."

# Calculation settings: formulas recalculate when opened in Excel or imported into Google Sheets.
try:
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
except AttributeError:
    pass


def fill(color):
    return PatternFill("solid", fgColor=color)


def set_cell(cell, value=None, *, font=None, fill_color=None, align=None, border=None, number_format=None):
    if value is not None:
        cell.value = value
    if font:
        cell.font = copy(font)
    if fill_color:
        cell.fill = fill(fill_color)
    if align:
        cell.alignment = copy(align)
    if border:
        cell.border = copy(border)
    if number_format:
        cell.number_format = number_format
    return cell


def style_title(sheet, title, subtitle, last_col):
    last_letter = get_column_letter(last_col)
    sheet.merge_cells(f"A1:{last_letter}2")
    c = sheet["A1"]
    c.value = title
    c.font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    c.fill = fill(NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 26
    sheet.row_dimensions[2].height = 26
    sheet.merge_cells(f"A3:{last_letter}3")
    c = sheet["A3"]
    c.value = subtitle
    c.font = Font(name="Aptos", size=10, italic=True, color=TEAL)
    c.fill = fill(PALE_TEAL)
    c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    sheet.row_dimensions[3].height = 26
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85


def style_section(sheet, row, start_col, end_col, text):
    start = get_column_letter(start_col)
    end = get_column_letter(end_col)
    sheet.merge_cells(f"{start}{row}:{end}{row}")
    c = sheet.cell(row, start_col)
    c.value = text
    c.font = Font(name="Aptos", size=11, bold=True, color=WHITE)
    c.fill = fill(TEAL)
    c.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[row].height = 22


def add_table(sheet, ref, name):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False
    )
    sheet.add_table(table)


def style_header_row(sheet, row, headers):
    for col, header in enumerate(headers, 1):
        c = sheet.cell(row, col)
        c.value = header
        c.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        c.fill = fill(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(top=medium_navy, bottom=medium_navy)
    sheet.row_dimensions[row].height = 34


def style_data_rows(sheet, start_row, end_row, ncols, formula_cols=None):
    formula_cols = set(formula_cols or [])
    for row in range(start_row, end_row + 1):
        for col in range(1, ncols + 1):
            c = sheet.cell(row, col)
            c.font = Font(name="Aptos", size=10, color=INK)
            c.fill = fill(PALE_BLUE if col in formula_cols else INPUT)
            c.alignment = Alignment(vertical="top", wrap_text=True)
            c.border = Border(bottom=thin_gray)
        sheet.row_dimensions[row].height = 21


def apply_number_formats(sheet, headers, start_row, end_row):
    for col, header in enumerate(headers, 1):
        h = header.lower()
        fmt = None
        if "%" in header:
            fmt = pct_fmt
        elif "date" in h or h in {"start", "end", "due", "completion", "follow-up date"}:
            fmt = date_fmt
        elif any(word in h for word in ["cost", "amount", "budget", "spend", "price", "quote", "balance", "paid", "difference", "fee", "discount"]):
            fmt = money_fmt
        elif any(word in h for word in ["quantity", "square footage", "width", "height", "depth", "length", "ceiling"]):
            fmt = '#,##0.00'
        if fmt:
            for row in range(start_row, end_row + 1):
                sheet.cell(row, col).number_format = fmt


def add_status_conditional_formatting(sheet, ranges):
    # Ranges are full A1 ranges; rules use the first cell in each range.
    for range_ref in ranges:
        first = range_ref.split(":")[0]
        col = "".join(ch for ch in first if ch.isalpha())
        row = "".join(ch for ch in first if ch.isdigit())
        cell_ref = f"{col}{row}"
        sheet.conditional_formatting.add(
            range_ref,
            FormulaRule(formula=[f'=OR({cell_ref}="Complete",{cell_ref}="Approved",{cell_ref}="Delivered",{cell_ref}="Installed",{cell_ref}="Good condition",{cell_ref}="Paid",{cell_ref}="On budget")'], fill=fill(GREEN))
        )
        sheet.conditional_formatting.add(
            range_ref,
            FormulaRule(formula=[f'=OR({cell_ref}="Pending",{cell_ref}="In progress",{cell_ref}="Sent",{cell_ref}="Waiting on client",{cell_ref}="Waiting on vendor",{cell_ref}="On hold",{cell_ref}="Partially paid",{cell_ref}="Scheduled")'], fill=fill(YELLOW))
        )
        sheet.conditional_formatting.add(
            range_ref,
            FormulaRule(formula=[f'=OR({cell_ref}="Overdue",{cell_ref}="Delayed",{cell_ref}="Damaged",{cell_ref}="Missing item",{cell_ref}="Rejected",{cell_ref}="Revisions requested")'], fill=fill(RED))
        )
        sheet.conditional_formatting.add(
            range_ref,
            FormulaRule(formula=[f'=OR({cell_ref}="Cancelled",{cell_ref}="On hold",{cell_ref}="Returned")'], fill=fill(GRAY))
        )


def add_validation(sheet, named_range, cell_range, prompt_title="Choose from the dropdown", prompt="Select a value from the list."):
    dv = DataValidation(type="list", formula1=f"={named_range}", allow_blank=True)
    dv.error = "Please choose a value from the dropdown list."
    dv.errorTitle = "Invalid entry"
    dv.promptTitle = prompt_title
    dv.prompt = prompt
    dv.showInputMessage = True
    dv.showErrorMessage = True
    sheet.add_data_validation(dv)
    dv.add(cell_range)
    return dv


def add_date_validation(sheet, cell_range):
    dv = DataValidation(type="date", operator="between", formula1="DATE(1900,1,1)", formula2="DATE(2100,12,31)", allow_blank=True)
    dv.error = "Enter a valid date between 1/1/1900 and 12/31/2100."
    dv.errorTitle = "Invalid date"
    dv.promptTitle = "Enter a date"
    dv.prompt = "Use a date such as 6/30/2026."
    dv.showInputMessage = True
    dv.showErrorMessage = True
    sheet.add_data_validation(dv)
    dv.add(cell_range)


def set_widths(sheet, widths):
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width


def make_data_sheet(title, subtitle, headers, widths, table_name, rows=300, formulas=None, tab_color=TEAL):
    sheet = wb.create_sheet(title)
    sheet.sheet_properties.tabColor = tab_color
    style_title(sheet, title, subtitle, len(headers))
    sheet["A4"] = "Tip: enter data in the warm cream cells. Blue cells calculate automatically. Use the filter arrows in row 6 to focus the view."
    sheet["A4"].font = Font(name="Aptos", size=9, italic=True, color=TEAL)
    sheet.merge_cells(start_row=4, start_column=1, end_row=4, end_column=len(headers))
    sheet["A4"].alignment = Alignment(wrap_text=True)
    style_header_row(sheet, 6, headers)
    end_row = 6 + rows
    formula_cols = set((formulas or {}).keys())
    style_data_rows(sheet, 7, end_row, len(headers), formula_cols)
    for col, formula in (formulas or {}).items():
        for row in range(7, end_row + 1):
            sheet.cell(row, col).value = formula(row)
            sheet.cell(row, col).fill = fill(PALE_BLUE)
    apply_number_formats(sheet, headers, 7, end_row)
    set_widths(sheet, widths)
    add_table(sheet, f"A6:{get_column_letter(len(headers))}{end_row}", table_name)
    sheet.freeze_panes = "A7"
    sheet.auto_filter.ref = f"A6:{get_column_letter(len(headers))}{end_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:6"
    sheet.sheet_view.selection[0].activeCell = "A7"
    return sheet, end_row


# ---------------------- Lists and named ranges ----------------------
lists = {
    "ProjectTypes": ["Residential", "Commercial", "Office", "Retail", "Hospitality", "Staging"],
    "ProjectStatuses": ["Planning", "In progress", "On hold", "Complete", "Archived"],
    "YesNo": ["Yes", "No"],
    "BillingMethods": ["Flat fee", "Hourly", "Percentage", "Hybrid"],
    "RoomTypes": ["Living room", "Bedroom", "Kitchen", "Bathroom", "Dining room", "Entryway", "Office", "Nursery", "Outdoor", "Commercial space", "Other"],
    "DesignStatuses": ["Not started", "Concept", "Sourcing", "Client review", "Approved", "Complete", "On hold"],
    "PriorityLevels": ["Low", "Medium", "High", "Urgent"],
    "FFECategories": ["Sofa", "Chair", "Table", "Bed", "Rug", "Lighting", "Artwork", "Mirror", "Window treatment", "Hardware", "Plumbing fixture", "Appliance", "Decor", "Textile", "Cabinetry", "Flooring", "Wallcovering", "Other"],
    "OrderStatuses": ["Not ordered", "Sourcing", "Option selected", "Client review", "Approved", "Ordered", "Confirmed", "In production", "Shipped", "Delivered", "Installed", "Cancelled", "Backordered"],
    "ApprovalStatuses": ["Not sent", "Sent", "Approved", "Revisions requested", "Rejected", "On hold"],
    "DeliveryStatuses": ["Not shipped", "Shipped", "In transit", "Delivered", "Delayed", "Damaged", "Missing item", "Cancelled"],
    "InstallStatuses": ["Not scheduled", "Scheduled", "In progress", "Installed", "Punch list", "On hold"],
    "FixtureTypes": ["Lighting", "Faucet", "Sink", "Shower fixture", "Tub", "Toilet", "Cabinet hardware", "Door hardware", "Appliance", "Electrical", "Other"],
    "FinishTypes": ["Paint", "Wallpaper", "Flooring", "Tile", "Countertop", "Backsplash", "Cabinet finish", "Fabric", "Trim", "Stone", "Wood", "Metal", "Other"],
    "SurfaceApplications": ["Wall", "Ceiling", "Floor", "Cabinet", "Countertop", "Shower wall", "Backsplash", "Upholstery", "Drapery"],
    "PreferredOptions": ["Option 1", "Option 2", "Option 3", "Custom", "Not selected"],
    "VendorTypes": ["Furniture", "Lighting", "Fabric", "Flooring", "Tile", "Paint", "Wallpaper", "Decor", "Hardware", "Contractor", "Installer", "Workroom", "Custom furniture", "Other"],
    "POStatuses": ["Not ordered", "Ordered", "Confirmed", "In production", "Shipped", "Delivered", "Delayed", "Cancelled", "Returned"],
    "DamageChecks": ["Not checked", "Good condition", "Damaged", "Missing item", "Wrong item"],
    "TimelinePhases": ["Consultation", "Measurements", "Concept design", "Sourcing", "Client approval", "Purchasing", "Construction", "Delivery", "Installation", "Styling", "Final reveal", "Punch list"],
    "TaskStatuses": ["Not started", "In progress", "Waiting on client", "Waiting on vendor", "Complete", "Delayed"],
    "MeasurementTypes": ["Room", "Wall", "Window", "Door", "Cabinet", "Furniture space", "Rug area", "Lighting location"],
    "Trades": ["General contractor", "Painter", "Electrician", "Plumber", "Flooring installer", "Tile installer", "Wallpaper installer", "Carpenter", "Window treatment installer", "Furniture installer", "Cleaner", "Other"],
    "PaymentStatuses": ["Draft", "Sent", "Paid", "Partially paid", "Overdue", "Cancelled"],
    "PaymentMethods": ["ACH", "Check", "Credit card", "Cash", "Wire", "Other"],
    "MeetingTypes": ["Client consultation", "Vendor meeting", "Site visit", "Installation meeting", "Contractor meeting", "Final walkthrough"],
    "TaskCategories": ["Design", "Sourcing", "Client", "Vendor", "Purchasing", "Delivery", "Installation", "Administration", "Other"],
}

lists_sheet = wb.create_sheet("Lists & Helpers")
lists_sheet.sheet_state = "hidden"
lists_sheet.sheet_view.showGridLines = False
for col, (name, values) in enumerate(lists.items(), 1):
    lists_sheet.cell(1, col).value = name
    lists_sheet.cell(1, col).font = Font(bold=True, color=WHITE)
    lists_sheet.cell(1, col).fill = fill(NAVY)
    for row, value in enumerate(values, 2):
        lists_sheet.cell(row, col).value = value
    lists_sheet.column_dimensions[get_column_letter(col)].width = max(16, min(28, len(name) + 4))

for col, name in enumerate(lists, 1):
    letter = get_column_letter(col)
    end = len(lists[name]) + 1
    wb.defined_names.add(DefinedName(name, attr_text=f"'Lists & Helpers'!${letter}$2:${letter}${end}"))
# These ranges intentionally include blank rows so the user can add rooms and vendors without editing validation rules.
wb.defined_names.add(DefinedName("RoomNames", attr_text="'Room List - Area Tracker'!$A$7:$A$56"))
wb.defined_names.add(DefinedName("VendorNames", attr_text="'Vendor & Supplier Database'!$A$7:$A$106"))

# ---------------------- Welcome / Instructions ----------------------
welcome = wb["Welcome"]
welcome.sheet_properties.tabColor = GOLD
style_title(welcome, "INTERIOR DESIGN PROJECT HUB", "A polished, editable project tracker for interior designers • compatible with Microsoft Excel + Google Sheets", 10)
set_widths(welcome, [24, 27, 27, 27, 25, 25, 25, 25, 20, 20])
for r in range(1, 60):
    welcome.row_dimensions[r].height = 21
welcome.row_dimensions[1].height = 30
welcome.row_dimensions[2].height = 30

welcome.merge_cells("A5:J5")
welcome["A5"] = "Welcome! Start with the five steps below, then use the dashboard to see the whole project at a glance."
welcome["A5"].font = Font(size=12, bold=True, color=INK)
welcome["A5"].fill = fill(PALE_TEAL)
welcome["A5"].alignment = Alignment(wrap_text=True, vertical="center")
welcome.row_dimensions[5].height = 30

style_section(welcome, 7, 1, 10, "QUICK START GUIDE")
steps = [
    ("Step 1", "Add client/project details", "Complete the Client & Project Profile tab. The project name, budget, dates, and status flow into the dashboard."),
    ("Step 2", "Add rooms", "List each area in Room List - Area Tracker. Use the room type, design status, budget allocation, and priority dropdowns."),
    ("Step 3", "Add FF&E items", "Enter furniture, fixtures, and equipment in Master FF&E Schedule. Total cost and final cost calculate automatically."),
    ("Step 4", "Track budget, orders, approvals, and delivery", "Use the Budget, Purchase Order, Client Approval, Delivery, Timeline, and Task tabs as work progresses."),
    ("Step 5", "Use dashboard for overview", "Review KPIs and charts for budget, item status, delivery, rooms, approvals, and open orders."),
]
for i, (step, title, body) in enumerate(steps, 8):
    welcome.cell(i, 1).value = step
    welcome.cell(i, 1).font = Font(bold=True, color=TEAL)
    welcome.cell(i, 1).fill = fill(PALE_BLUE)
    welcome.cell(i, 2).value = title
    welcome.cell(i, 2).font = Font(bold=True, color=INK)
    welcome.cell(i, 2).fill = fill(INPUT)
    welcome.merge_cells(start_row=i, start_column=3, end_row=i, end_column=10)
    welcome.cell(i, 3).value = body
    welcome.cell(i, 3).alignment = Alignment(wrap_text=True, vertical="center")
    welcome.cell(i, 3).fill = fill(INPUT)
    for col in range(1, 11):
        welcome.cell(i, col).border = Border(bottom=thin_gray)
    welcome.row_dimensions[i].height = 32

style_section(welcome, 15, 1, 10, "WHAT EACH TAB IS USED FOR")
tab_guide = [
    ("Project Dashboard", "Main overview: KPI cards and charts for budget, items, approvals, delivery, and orders."),
    ("Client & Project Profile", "Project summary, client contact details, dates, budget, scope, notes, and status."),
    ("Room List - Area Tracker", "Room-by-room planning, budget allocation, estimated/actual spend, design status, and priorities."),
    ("Master FF&E Schedule", "Core Furniture, Fixtures & Equipment register with pricing, links, vendors, approvals, delivery, and install status."),
    ("Furniture Schedule", "Furniture-only view for quick sourcing and purchasing."),
    ("Fixtures Schedule", "Lighting, plumbing, hardware, appliances, and built-in fixtures."),
    ("Finishes Schedule", "Paint, wallpaper, flooring, tile, stone, fabric, and other finish selections."),
    ("Product Sourcing Tracker", "Compare up to three product options before selecting a final item."),
    ("Vendor & Supplier Database", "Central vendor contacts, trade terms, discounts, lead times, and policies."),
    ("Budget Tracker", "Category budget, estimated/actual spend, paid amount, and balance due."),
    ("Purchase Order - Order Tracker", "POs, order dates, delivery expectations, tracking, invoices, and paid status."),
    ("Delivery & Installation Tracker", "Expected/actual delivery, damage checks, storage, installers, and installation status."),
    ("Client Approval Tracker", "Items and decisions sent to clients, responses, revisions, comments, and final approval."),
    ("Project Timeline - Gantt", "Phases, tasks, owners, dates, dependencies, priority, and status."),
    ("Task & To-Do List", "Daily task list with room, category, owner, due date, priority, and status."),
    ("Measurement Log", "Room, wall, window, door, cabinet, rug, furniture-space, and lighting measurements."),
    ("Contractor - Installer Tracker", "Trades, scope, quotes, schedule, insurance, contracts, and payment status."),
    ("Invoice & Payment Tracker", "Client invoices, dates, amounts, due dates, paid dates, and payment method."),
    ("Project Notes - Meeting Notes", "Meeting attendees, discussion points, decisions, actions, and follow-up dates."),
    ("Printable Reports", "Client-facing FF&E, room budget, vendor order, installation, and final punch-list print views."),
]
welcome.append(["Tab", "Use it for"])
for c in welcome[16]:
    c.font = Font(bold=True, color=WHITE)
    c.fill = fill(NAVY)
    c.alignment = Alignment(horizontal="center")
for tab, desc in tab_guide:
    welcome.append([tab, desc])
    row = welcome.max_row
    welcome.cell(row, 1).font = Font(bold=True, color=TEAL)
    welcome.cell(row, 2).alignment = Alignment(wrap_text=True)
    for col in range(1, 11):
        welcome.cell(row, col).border = Border(bottom=thin_gray)
        if col > 2:
            welcome.cell(row, col).fill = fill(WHITE)
    welcome.merge_cells(start_row=row, start_column=2, end_row=row, end_column=10)
    welcome.row_dimensions[row].height = 28

last_guide_row = welcome.max_row
style_section(welcome, last_guide_row + 2, 1, 5, "HOW TO ENTER DATA")
entry_start = last_guide_row + 3
entry_notes = [
    ("Warm cream cells", "Editable input cells. Click a cell and type; use dropdown arrows wherever they appear."),
    ("Blue cells", "Formula cells. They calculate automatically; you can still edit the workbook because this is an unlocked template."),
    ("Dropdowns", "Choose a value from the dropdown for consistent reporting. Add rooms and vendors in their dedicated tabs first; those lists power the dropdowns."),
    ("Links", "Paste full product, image, website, tracking, or invoice links. Keep links in the dedicated link columns."),
]
for idx, (label, body) in enumerate(entry_notes, entry_start):
    welcome.cell(idx, 1).value = label
    welcome.cell(idx, 1).font = Font(bold=True, color=INK)
    welcome.cell(idx, 1).fill = fill(INPUT if idx != entry_start + 1 else PALE_BLUE)
    welcome.merge_cells(start_row=idx, start_column=2, end_row=idx, end_column=10)
    welcome.cell(idx, 2).value = body
    welcome.cell(idx, 2).alignment = Alignment(wrap_text=True)
    welcome.cell(idx, 2).fill = fill(WHITE)
    welcome.row_dimensions[idx].height = 28

style_section(welcome, entry_start + 5, 1, 5, "COLOR CODING")
color_row = entry_start + 6
color_items = [("Green", GREEN, "Complete / approved / delivered / installed / paid / on budget"), ("Yellow", YELLOW, "Pending / in progress / waiting / scheduled"), ("Red", RED, "Overdue / over budget / delayed / damaged / rejected"), ("Gray", GRAY, "Cancelled / returned / on hold")]
for col, (label, color, desc) in enumerate(color_items, 1):
    welcome.cell(color_row, col).value = label
    welcome.cell(color_row, col).font = Font(bold=True, color=INK)
    welcome.cell(color_row, col).fill = fill(color)
    welcome.cell(color_row + 1, col).value = desc
    welcome.cell(color_row + 1, col).alignment = Alignment(wrap_text=True, vertical="top")
    welcome.cell(color_row + 1, col).fill = fill(WHITE)
    welcome.column_dimensions[get_column_letter(col)].width = max(welcome.column_dimensions[get_column_letter(col)].width or 0, 25)
welcome.merge_cells(start_row=color_row + 1, start_column=1, end_row=color_row + 1, end_column=2)
welcome.merge_cells(start_row=color_row + 1, start_column=3, end_row=color_row + 1, end_column=4)
welcome.merge_cells(start_row=color_row + 1, start_column=5, end_row=color_row + 1, end_column=6)
welcome.merge_cells(start_row=color_row + 1, start_column=7, end_row=color_row + 1, end_column=8)
welcome.row_dimensions[color_row + 1].height = 38

style_section(welcome, color_row + 3, 1, 10, "COMPATIBILITY, EDITING & SUPPORT")
compat_row = color_row + 4
welcome.merge_cells(start_row=compat_row, start_column=1, end_row=compat_row, end_column=10)
welcome.cell(compat_row, 1).value = "Compatibility note: this workbook is designed for Microsoft Excel and Google Sheets. Upload the .xlsx file to Google Drive and choose Open with Google Sheets. Standard formulas, filters, dropdowns, conditional formatting, and charts are used for broad compatibility."
welcome.cell(compat_row, 1).alignment = Alignment(wrap_text=True, vertical="center")
welcome.cell(compat_row, 1).fill = fill(PALE_TEAL)
welcome.row_dimensions[compat_row].height = 42
welcome.merge_cells(start_row=compat_row + 1, start_column=1, end_row=compat_row + 1, end_column=10)
welcome.cell(compat_row + 1, 1).value = "Editable / unlocked version: no sheet protection is applied. Update labels, formulas, colors, dropdown source lists, and layouts to fit your studio workflow. Keep formula cells intact if you want the dashboard to stay accurate."
welcome.cell(compat_row + 1, 1).alignment = Alignment(wrap_text=True, vertical="center")
welcome.cell(compat_row + 1, 1).fill = fill(INPUT)
welcome.row_dimensions[compat_row + 1].height = 42
welcome.merge_cells(start_row=compat_row + 2, start_column=1, end_row=compat_row + 2, end_column=10)
welcome.cell(compat_row + 2, 1).value = "Support / contact: replace this line with your studio name, support email, website, or template help desk before sharing with clients."
welcome.cell(compat_row + 2, 1).alignment = Alignment(wrap_text=True, vertical="center")
welcome.cell(compat_row + 2, 1).fill = fill(WHITE)
welcome.row_dimensions[compat_row + 2].height = 30
welcome.freeze_panes = "A7"
welcome.page_setup.orientation = "landscape"
welcome.page_setup.fitToWidth = 1
welcome.page_setup.fitToHeight = 0
welcome.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------- Project Profile ----------------------
profile = wb.create_sheet("Client & Project Profile")
profile.sheet_properties.tabColor = TEAL
style_title(profile, "CLIENT & PROJECT PROFILE", "Complete this clean project summary first. Key fields feed the Project Dashboard.", 10)
set_widths(profile, [25, 30, 4, 25, 30, 18, 18, 18, 18, 18])
style_section(profile, 5, 1, 2, "CLIENT & CONTACT")
style_section(profile, 5, 4, 10, "PROJECT SUMMARY")
profile_fields = [
    (6, "Client name", ""), (7, "Client email", ""), (8, "Phone number", ""), (9, "Project name", ""),
    (10, "Project address", ""), (11, "Project type", ""), (12, "Design style", ""), (13, "Start date", ""),
    (14, "Target completion date", ""), (15, "Total budget", ""), (16, "Designer name", ""), (17, "Project status", ""),
    (18, "Contract signed", ""), (19, "Deposit paid", ""), (20, "Design package selected", ""), (21, "Billing method", ""),
]
for row, label, value in profile_fields:
    profile.cell(row, 1).value = label
    profile.cell(row, 1).font = Font(bold=True, color=INK)
    profile.cell(row, 1).fill = fill(PALE_TEAL)
    profile.cell(row, 2).value = value
    profile.cell(row, 2).fill = fill(INPUT)
    profile.cell(row, 2).border = Border(bottom=thin_gray)
    profile.cell(row, 2).alignment = Alignment(wrap_text=True, vertical="top")
    profile.row_dimensions[row].height = 24
for row, label in [(6, "Client notes"), (10, "Scope summary"), (15, "Project status notes")]:
    profile.cell(row, 4).value = label
    profile.cell(row, 4).font = Font(bold=True, color=INK)
    profile.cell(row, 4).fill = fill(PALE_TEAL)
    profile.merge_cells(start_row=row, start_column=5, end_row=row + (2 if row != 15 else 1), end_column=10)
    profile.cell(row, 5).fill = fill(INPUT)
    profile.cell(row, 5).alignment = Alignment(wrap_text=True, vertical="top")
    for rr in range(row, row + (3 if row != 15 else 2)):
        for cc in range(5, 11):
            profile.cell(rr, cc).border = Border(bottom=thin_gray)
    profile.row_dimensions[row].height = 26
for r in [13, 14]:
    profile.cell(r, 2).number_format = date_fmt
profile.cell(15, 2).number_format = whole_money_fmt
add_validation(profile, "ProjectTypes", "B11")
add_validation(profile, "ProjectStatuses", "B17")
add_validation(profile, "YesNo", "B18:B19")
add_validation(profile, "BillingMethods", "B21")
add_date_validation(profile, "B13:B14")
profile.freeze_panes = "A6"
profile.page_setup.orientation = "landscape"
profile.page_setup.fitToWidth = 1
profile.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------- Room List ----------------------
room_headers = ["Room name", "Room type", "Floor/level", "Square footage", "Design status", "Budget allocation", "Estimated spend", "Actual spend", "Remaining budget", "Priority level", "Notes"]
room_widths = [24, 19, 14, 15, 18, 17, 17, 17, 18, 14, 35]
rooms, room_end = make_data_sheet("Room List - Area Tracker", "Organize the project by room or area. Remaining budget calculates as allocation less actual spend.", room_headers, room_widths, "RoomAreaTracker", 50, {9: lambda r: f'=IF(A{r}="","",F{r}-H{r})'}, GOLD)
add_validation(rooms, "RoomTypes", "B7:B56")
add_validation(rooms, "DesignStatuses", "E7:E56")
add_validation(rooms, "PriorityLevels", "J7:J56")
add_status_conditional_formatting(rooms, ["E7:E56", "J7:J56"])
rooms.conditional_formatting.add("I7:I56", FormulaRule(formula=['=AND($A7<>"",$I7<0)'], fill=fill(RED)))

# ---------------------- Master FF&E ----------------------
ffe_headers = ["Item ID", "Room", "Category", "Item name", "Description", "Quantity", "Dimensions", "Material", "Finish", "Color", "Vendor", "Brand", "Product link", "SKU/model number", "Unit cost", "Total cost", "Tax", "Shipping", "Discount", "Final cost", "Lead time", "Order status", "Approval status", "Delivery status", "Install status", "Notes", "Image link"]
ffe_widths = [12, 20, 19, 27, 32, 10, 17, 17, 17, 14, 22, 18, 30, 20, 14, 15, 13, 13, 13, 15, 14, 17, 18, 17, 16, 32, 30]
ffe, ffe_end = make_data_sheet("Master FF&E Schedule", "Core Furniture, Fixtures & Equipment schedule. Enter quantity and unit cost; total cost and final cost calculate automatically.", ffe_headers, ffe_widths, "MasterFFE", 300, {16: lambda r: f'=IF(D{r}="","",F{r}*O{r})', 20: lambda r: f'=IF(D{r}="","",P{r}+Q{r}+R{r}-S{r})'}, TEAL)
add_validation(ffe, "RoomNames", "B7:B306")
add_validation(ffe, "FFECategories", "C7:C306")
add_validation(ffe, "VendorNames", "K7:K306")
add_validation(ffe, "OrderStatuses", "V7:V306")
add_validation(ffe, "ApprovalStatuses", "W7:W306")
add_validation(ffe, "DeliveryStatuses", "X7:X306")
add_validation(ffe, "InstallStatuses", "Y7:Y306")
add_status_conditional_formatting(ffe, ["V7:V306", "W7:W306", "X7:X306", "Y7:Y306"])

# ---------------------- Furniture ----------------------
furniture_headers = ["Room", "Furniture item", "Dimensions", "Quantity", "Vendor", "Product link", "Fabric/material", "Finish", "Cost", "Approval status", "Order status", "Delivery date", "Install date"]
furniture_widths = [20, 28, 18, 10, 22, 30, 20, 17, 14, 18, 17, 15, 15]
furniture, furniture_end = make_data_sheet("Furniture Schedule", "A focused furniture-only list for quick sourcing, approval, and purchasing.", furniture_headers, furniture_widths, "FurnitureSchedule", 150, {}, "7F9E9F")
add_validation(furniture, "RoomNames", "A7:A156")
add_validation(furniture, "VendorNames", "E7:E156")
add_validation(furniture, "ApprovalStatuses", "J7:J156")
add_validation(furniture, "POStatuses", "K7:K156")
add_date_validation(furniture, "L7:M156")
add_status_conditional_formatting(furniture, ["J7:J156", "K7:K156"])

# ---------------------- Fixtures ----------------------
fixture_headers = ["Room", "Fixture type", "Item name", "Finish", "Size", "Vendor", "Model/SKU", "Quantity", "Unit cost", "Total cost", "Required rough-in date", "Order deadline", "Delivery date", "Installer notes"]
fixture_widths = [20, 20, 26, 18, 15, 22, 20, 10, 14, 15, 20, 17, 15, 34]
fixtures, fixture_end = make_data_sheet("Fixtures Schedule", "Lighting, plumbing, hardware, appliances, and built-in fixtures.", fixture_headers, fixture_widths, "FixturesSchedule", 150, {10: lambda r: f'=IF(C{r}="","",H{r}*I{r})'}, "7F9E9F")
add_validation(fixtures, "RoomNames", "A7:A156")
add_validation(fixtures, "FixtureTypes", "B7:B156")
add_validation(fixtures, "VendorNames", "F7:F156")
add_date_validation(fixtures, "K7:M156")

# ---------------------- Finishes ----------------------
finish_headers = ["Room", "Surface/application", "Finish type", "Product name", "Color/finish", "Material", "Vendor", "Product link", "SKU", "Quantity needed", "Unit cost", "Total cost", "Sample ordered", "Sample approved", "Installer notes"]
finish_widths = [20, 21, 19, 27, 20, 18, 22, 30, 18, 15, 14, 15, 16, 16, 34]
finishes, finish_end = make_data_sheet("Finishes Schedule", "Track finish selections, samples, quantities, pricing, and installer notes by room.", finish_headers, finish_widths, "FinishesSchedule", 150, {12: lambda r: f'=IF(D{r}="","",J{r}*K{r})'}, "7F9E9F")
add_validation(finishes, "RoomNames", "A7:A156")
add_validation(finishes, "SurfaceApplications", "B7:B156")
add_validation(finishes, "FinishTypes", "C7:C156")
add_validation(finishes, "VendorNames", "G7:G156")
add_validation(finishes, "YesNo", "M7:N156")

# ---------------------- Product Sourcing ----------------------
sourcing_headers = ["Room", "Item needed", "Option 1 link", "Option 1 price", "Option 2 link", "Option 2 price", "Option 3 link", "Option 3 price", "Preferred option", "Pros", "Cons", "Client feedback", "Final selection"]
sourcing_widths = [20, 26, 30, 14, 30, 14, 30, 14, 17, 28, 28, 30, 28]
sourcing, sourcing_end = make_data_sheet("Product Sourcing Tracker", "Compare up to three options before committing to a final FF&E selection.", sourcing_headers, sourcing_widths, "ProductSourcing", 150, {}, "D6A756")
add_validation(sourcing, "RoomNames", "A7:A156")
add_validation(sourcing, "PreferredOptions", "I7:I156")

# ---------------------- Vendors ----------------------
vendor_headers = ["Vendor name", "Vendor type", "Contact person", "Email", "Phone", "Website", "Account number", "Trade discount %", "Lead time", "Shipping terms", "Return policy", "Notes"]
vendor_widths = [24, 21, 22, 27, 17, 30, 18, 17, 15, 25, 25, 34]
vendors, vendor_end = make_data_sheet("Vendor & Supplier Database", "Create your vendor list first; vendor names then appear in dropdowns throughout the workbook.", vendor_headers, vendor_widths, "VendorDatabase", 100, {}, GOLD)
add_validation(vendors, "VendorTypes", "B7:B106")

# ---------------------- Budget ----------------------
budget_headers = ["Budget category", "Budgeted amount", "Estimated amount", "Actual amount", "Difference", "Paid amount", "Balance due", "Payment status", "Notes"]
budget_widths = [25, 17, 17, 17, 18, 16, 16, 18, 34]
budget, budget_end = make_data_sheet("Budget Tracker", "Track category budgets, estimates, actuals, payments, and balance due. Difference = budgeted less estimated.", budget_headers, budget_widths, "BudgetTracker", 300, {5: lambda r: f'=IF(A{r}="","",B{r}-C{r})', 7: lambda r: f'=IF(A{r}="","",D{r}-F{r})'}, TEAL)
budget_categories = ["Furniture", "Lighting", "Rugs", "Decor", "Artwork", "Window treatments", "Flooring", "Paint/wallcovering", "Contractor labor", "Installation", "Shipping", "Taxes", "Design fee", "Contingency"]
# Add budget categories to a dedicated named range rather than mixing them with product categories.
helper_col = len(lists) + 1
lists_sheet.cell(1, helper_col).value = "BudgetCategories"
for row, value in enumerate(budget_categories, 2):
    lists_sheet.cell(row, helper_col).value = value
wb.defined_names.add(DefinedName("BudgetCategories", attr_text=f"'Lists & Helpers'!${get_column_letter(helper_col)}$2:${get_column_letter(helper_col)}${len(budget_categories)+1}"))
add_validation(budget, "BudgetCategories", "A7:A306")
add_validation(budget, "PaymentStatuses", "H7:H306")
add_status_conditional_formatting(budget, ["H7:H306"])
# Totals block keeps the budget tab useful on its own, in addition to the dashboard KPIs.
budget_total_row = budget_end + 2
budget.cell(budget_total_row, 1).value = "TOTAL PROJECT BUDGET"
budget.cell(budget_total_row, 2).value = "=SUM(B7:B306)"
budget.cell(budget_total_row, 3).value = "=SUM(C7:C306)"
budget.cell(budget_total_row, 4).value = "=SUM(D7:D306)"
budget.cell(budget_total_row, 5).value = f"=B{budget_total_row}-C{budget_total_row}"
budget.cell(budget_total_row, 6).value = "=SUM(F7:F306)"
budget.cell(budget_total_row, 7).value = f"=D{budget_total_row}-F{budget_total_row}"
budget.cell(budget_total_row, 8).value = '=COUNTIF(H7:H306,"Paid")&" paid / "&COUNTIFS(A7:A306,"<>",H7:H306,"<>Paid")&" unpaid"'
budget.cell(budget_total_row, 9).value = "Auto-calculated totals"
for col in range(1, 10):
    c = budget.cell(budget_total_row, col)
    c.font = Font(name="Aptos", size=10, bold=True, color=NAVY)
    c.fill = fill(PALE_BLUE)
    c.border = Border(top=medium_navy, bottom=medium_navy)
    c.alignment = Alignment(vertical="center", wrap_text=True)
    if 2 <= col <= 7:
        c.number_format = money_fmt
budget.row_dimensions[budget_total_row].height = 28
budget.conditional_formatting.add("E7:E306", FormulaRule(formula=['=AND($A7<>"",$E7<0)'], fill=fill(RED)))
budget.conditional_formatting.add("G7:G306", FormulaRule(formula=['=AND($A7<>"",$G7>0)'], fill=fill(YELLOW)))

# ---------------------- Orders ----------------------
order_headers = ["PO number", "Vendor", "Room", "Item", "Quantity", "Order date", "Expected ship date", "Expected delivery date", "Actual delivery date", "Order status", "Tracking number", "Carrier", "Invoice number", "Amount", "Paid status", "Notes"]
order_widths = [15, 22, 20, 28, 10, 14, 18, 21, 19, 17, 23, 16, 17, 15, 16, 32]
orders, orders_end = make_data_sheet("Purchase Order - Order Tracker", "Track every vendor order from PO creation through delivery, tracking, invoicing, and payment.", order_headers, order_widths, "PurchaseOrderTracker", 300, {}, TEAL)
add_validation(orders, "VendorNames", "B7:B306")
add_validation(orders, "RoomNames", "C7:C306")
add_validation(orders, "POStatuses", "J7:J306")
add_validation(orders, "YesNo", "O7:O306")
add_date_validation(orders, "F7:I306")
add_status_conditional_formatting(orders, ["J7:J306", "O7:O306"])
# Expected delivery date becomes red when overdue and the order is not finished.
orders.conditional_formatting.add("H7:H306", FormulaRule(formula=['=AND($H7<TODAY(),$H7<>"",$J7<>"Delivered",$J7<>"Cancelled",$A7<>"")'], fill=fill(RED)))

# ---------------------- Delivery ----------------------
delivery_headers = ["Item", "Room", "Vendor", "Delivery address", "Expected delivery date", "Actual delivery date", "Delivery status", "Received by", "Damage check", "Damage notes", "Storage location", "Install date", "Installer", "Installation status", "Punch list notes"]
delivery_widths = [28, 20, 22, 30, 21, 19, 17, 20, 17, 30, 22, 15, 22, 19, 34]
delivery, delivery_end = make_data_sheet("Delivery & Installation Tracker", "Keep deliveries organized: expected dates, damage checks, storage, installers, and installation status.", delivery_headers, delivery_widths, "DeliveryInstallationTracker", 300, {}, GOLD)
add_validation(delivery, "RoomNames", "B7:B306")
add_validation(delivery, "VendorNames", "C7:C306")
add_validation(delivery, "DeliveryStatuses", "G7:G306")
add_validation(delivery, "DamageChecks", "I7:I306")
add_validation(delivery, "InstallStatuses", "N7:N306")
add_date_validation(delivery, "E7:F306")
add_date_validation(delivery, "L7:L306")
add_status_conditional_formatting(delivery, ["G7:G306", "I7:I306", "N7:N306"])

# ---------------------- Approvals ----------------------
approval_headers = ["Item/design decision", "Room", "Approval type", "Sent to client date", "Client response date", "Status", "Approved option", "Client comments", "Revision needed", "Final approval date"]
approval_widths = [30, 20, 20, 20, 21, 19, 25, 34, 16, 20]
approvals, approvals_end = make_data_sheet("Client Approval Tracker", "Make client decisions visible: send, respond, revise, and record final approval dates.", approval_headers, approval_widths, "ClientApprovalTracker", 200, {}, TEAL)
add_validation(approvals, "RoomNames", "B7:B206")
add_validation(approvals, "ApprovalStatuses", "F7:F206")
add_validation(approvals, "YesNo", "I7:I206")
# Approval type choices are stored on the helper sheet for a dedicated dropdown.
approval_types = ["Furniture", "Finish", "Layout", "Color palette", "Budget", "Purchase", "Vendor", "Design concept"]
approval_col = helper_col + 1
lists_sheet.cell(1, approval_col).value = "ApprovalTypes"
for row, value in enumerate(approval_types, 2):
    lists_sheet.cell(row, approval_col).value = value
wb.defined_names.add(DefinedName("ApprovalTypes", attr_text=f"'Lists & Helpers'!${get_column_letter(approval_col)}$2:${get_column_letter(approval_col)}${len(approval_types)+1}"))
add_validation(approvals, "ApprovalTypes", "C7:C206")
add_date_validation(approvals, "D7:E206")
add_date_validation(approvals, "J7:J206")
add_status_conditional_formatting(approvals, ["F7:F206", "I7:I206"])

# ---------------------- Timeline ----------------------
timeline_headers = ["Phase", "Task", "Assigned to", "Start date", "Due date", "Completion date", "Status", "Priority", "Dependencies", "Notes"]
timeline_widths = [20, 30, 22, 15, 15, 20, 18, 14, 24, 34]
timeline, timeline_end = make_data_sheet("Project Timeline - Gantt", "Manage phases, tasks, owners, dependencies, dates, and priorities. Use the dates for a simple Gantt-style workflow.", timeline_headers, timeline_widths, "ProjectTimeline", 200, {}, GOLD)
add_validation(timeline, "TimelinePhases", "A7:A206")
add_validation(timeline, "TaskStatuses", "G7:G206")
add_validation(timeline, "PriorityLevels", "H7:H206")
add_date_validation(timeline, "D7:F206")
add_status_conditional_formatting(timeline, ["G7:G206", "H7:H206"])
timeline.conditional_formatting.add("E7:E206", FormulaRule(formula=['=AND($E7<TODAY(),$E7<>"",$G7<>"Complete",$B7<>"")'], fill=fill(RED)))

# ---------------------- Tasks ----------------------
task_headers = ["Task", "Room", "Category", "Assigned to", "Priority", "Due date", "Status", "Notes"]
task_widths = [32, 20, 19, 22, 14, 15, 20, 38]
tasks, tasks_end = make_data_sheet("Task & To-Do List", "Use this as the daily action list. Overdue open tasks are counted on the dashboard.", task_headers, task_widths, "TaskTodoList", 300, {}, TEAL)
add_validation(tasks, "RoomNames", "B7:B306")
add_validation(tasks, "TaskCategories", "C7:C306")
add_validation(tasks, "PriorityLevels", "E7:E306")
add_validation(tasks, "TaskStatuses", "G7:G306")
add_date_validation(tasks, "F7:F306")
add_status_conditional_formatting(tasks, ["E7:E306", "G7:G306"])
tasks.conditional_formatting.add("F7:F306", FormulaRule(formula=['=AND($F7<TODAY(),$F7<>"",$G7<>"Complete",$A7<>"")'], fill=fill(RED)))

# ---------------------- Measurements ----------------------
measurement_headers = ["Room", "Measurement type", "Width", "Height", "Depth", "Length", "Ceiling height", "Window size", "Door size", "Wall length", "Notes"]
measurement_widths = [20, 22, 14, 14, 14, 14, 17, 18, 17, 17, 36]
measurements, measurements_end = make_data_sheet("Measurement Log", "Record dimensions for rooms, walls, openings, furniture spaces, rugs, and lighting locations.", measurement_headers, measurement_widths, "MeasurementLog", 200, {}, "7F9E9F")
add_validation(measurements, "RoomNames", "A7:A206")
add_validation(measurements, "MeasurementTypes", "B7:B206")

# ---------------------- Contractors ----------------------
contractor_headers = ["Contractor/vendor name", "Trade", "Contact", "Scope", "Quote amount", "Start date", "End date", "Payment status", "Insurance received", "Contract signed", "Notes"]
contractor_widths = [27, 24, 24, 34, 16, 15, 15, 18, 18, 16, 34]
contractors, contractors_end = make_data_sheet("Contractor - Installer Tracker", "Coordinate construction and installation partners, scope, quotes, dates, documents, and payments.", contractor_headers, contractor_widths, "ContractorInstallerTracker", 150, {}, "7F9E9F")
add_validation(contractors, "Trades", "B7:B156")
add_validation(contractors, "PaymentStatuses", "H7:H156")
add_validation(contractors, "YesNo", "I7:J156")
add_date_validation(contractors, "F7:G156")
add_status_conditional_formatting(contractors, ["H7:H156", "I7:I156", "J7:J156"])

# ---------------------- Invoices ----------------------
invoice_headers = ["Invoice number", "Invoice date", "Client", "Description", "Amount", "Due date", "Paid date", "Payment status", "Payment method", "Notes"]
invoice_widths = [17, 15, 25, 34, 15, 15, 15, 19, 18, 34]
invoices, invoices_end = make_data_sheet("Invoice & Payment Tracker", "Track client billing, due dates, paid dates, status, and payment method.", invoice_headers, invoice_widths, "InvoicePaymentTracker", 200, {}, GOLD)
add_validation(invoices, "PaymentStatuses", "H7:H206")
add_validation(invoices, "PaymentMethods", "I7:I206")
add_date_validation(invoices, "B7:B206")
add_date_validation(invoices, "F7:G206")
add_status_conditional_formatting(invoices, ["H7:H206"])
invoices.conditional_formatting.add("F7:F206", FormulaRule(formula=['=AND($F7<TODAY(),$F7<>"",$H7<>"Paid",$H7<>"Cancelled",$A7<>"")'], fill=fill(RED)))

# ---------------------- Notes ----------------------
notes_headers = ["Date", "Meeting type", "Attendees", "Discussion points", "Decisions made", "Action items", "Follow-up date", "Notes"]
notes_widths = [15, 24, 28, 34, 34, 34, 18, 34]
notes, notes_end = make_data_sheet("Project Notes - Meeting Notes", "Capture conversations, decisions, action items, and follow-up dates in one searchable log.", notes_headers, notes_widths, "ProjectMeetingNotes", 200, {}, TEAL)
add_validation(notes, "MeetingTypes", "B7:B206")
add_date_validation(notes, "A7:A206")
add_date_validation(notes, "G7:G206")

# ---------------------- Dashboard helper data ----------------------
helper = wb.create_sheet("Dashboard Data")
helper.sheet_state = "hidden"
helper.sheet_view.showGridLines = False
# Room data for budget by room.
helper["A1"] = "Room"
helper["B1"] = "Budget allocation"
for r in range(2, 52):
    source = r + 5
    helper.cell(r, 1).value = f"='Room List - Area Tracker'!A{source}"
    helper.cell(r, 2).value = f"=IFERROR('Room List - Area Tracker'!F{source},0)"
# Budget vs actual summary.
helper["D1"] = "Metric"
helper["E1"] = "Amount"
for idx, metric in enumerate(["Budget", "Estimated", "Actual"], 2):
    helper.cell(idx, 4).value = metric
helper["E2"] = "=IFERROR(IF('Client & Project Profile'!$B$15>0,'Client & Project Profile'!$B$15,SUM('Budget Tracker'!$B$7:$B$306)),0)"
helper["E3"] = "=SUM('Budget Tracker'!$C$7:$C$306)"
helper["E4"] = "=SUM('Budget Tracker'!$D$7:$D$306)"
# Item status breakdown from Master FF&E order status.
helper["G1"] = "Item status"
helper["H1"] = "Items"
for idx, status in enumerate(lists["OrderStatuses"], 2):
    helper.cell(idx, 7).value = status
    helper.cell(idx, 8).value = f'=COUNTIF(\'Master FF&E Schedule\'!$V$7:$V$306,G{idx})'
# Delivery status breakdown.
helper["J1"] = "Delivery status"
helper["K1"] = "Items"
for idx, status in enumerate(lists["DeliveryStatuses"], 2):
    helper.cell(idx, 10).value = status
    helper.cell(idx, 11).value = f'=COUNTIF(\'Master FF&E Schedule\'!$X$7:$X$306,J{idx})'
# Category spend breakdown.
helper["M1"] = "Category"
helper["N1"] = "Final cost"
for idx, category in enumerate(lists["FFECategories"], 2):
    helper.cell(idx, 13).value = category
    helper.cell(idx, 14).value = f'=SUMIF(\'Master FF&E Schedule\'!$C$7:$C$306,M{idx},\'Master FF&E Schedule\'!$T$7:$T$306)'
for row in range(1, 40):
    for col in range(1, 15):
        helper.cell(row, col).font = Font(size=9)

# ---------------------- Dashboard ----------------------
dash = wb.create_sheet("Project Dashboard", 1)
dash.sheet_properties.tabColor = NAVY
style_title(dash, "PROJECT DASHBOARD", "A decision-ready overview of budget, scope, approvals, orders, deliveries, and installation.", 18)
set_widths(dash, [15, 15, 15, 15, 15, 15, 15, 15, 15, 15, 15, 15, 15, 15, 15, 15, 15, 15])
dash.merge_cells("A4:R4")
dash["A4"] = "Tip: enter data in the warm cream cells on the project tabs. Dashboard formulas recalculate automatically."
dash["A4"].font = Font(size=9, italic=True, color=TEAL)
dash["A4"].alignment = Alignment(wrap_text=True)

kpis = [
    ("Total project budget", "=IFERROR(IF('Client & Project Profile'!$B$15>0,'Client & Project Profile'!$B$15,SUM('Budget Tracker'!$B$7:$B$306)),0)", whole_money_fmt),
    ("Total estimated cost", "=SUM('Budget Tracker'!$C$7:$C$306)", whole_money_fmt),
    ("Total actual cost", "=SUM('Budget Tracker'!$D$7:$D$306)", whole_money_fmt),
    ("Remaining budget", "=A6-G6", whole_money_fmt),
    ("Budget used %", "=IFERROR(G6/A6,0)", pct_fmt),
    ("Number of rooms", "=COUNTIF('Room List - Area Tracker'!$A$7:$A$56,\"<>\")", '0'),
    ("Total FF&E items", "=COUNTIF('Master FF&E Schedule'!$D$7:$D$306,\"<>\")", '0'),
    ("Items approved", "=COUNTIF('Master FF&E Schedule'!$W$7:$W$306,\"Approved\")", '0'),
    ("Items ordered", "=COUNTIF('Master FF&E Schedule'!$V$7:$V$306,\"Ordered\")+COUNTIF('Master FF&E Schedule'!$V$7:$V$306,\"Confirmed\")+COUNTIF('Master FF&E Schedule'!$V$7:$V$306,\"In production\")+COUNTIF('Master FF&E Schedule'!$V$7:$V$306,\"Shipped\")+COUNTIF('Master FF&E Schedule'!$V$7:$V$306,\"Delivered\")+COUNTIF('Master FF&E Schedule'!$V$7:$V$306,\"Installed\")", '0'),
    ("Items delivered", "=COUNTIF('Master FF&E Schedule'!$X$7:$X$306,\"Delivered\")", '0'),
    ("Items installed", "=COUNTIF('Master FF&E Schedule'!$Y$7:$Y$306,\"Installed\")", '0'),
    ("Overdue tasks", "=COUNTIFS('Task & To-Do List'!$F$7:$F$306,\"<\"&TODAY(),'Task & To-Do List'!$F$7:$F$306,\"<>\",'Task & To-Do List'!$G$7:$G$306,\"<>Complete\",'Task & To-Do List'!$A$7:$A$306,\"<>\")", '0'),
    ("Pending client approvals", "=COUNTIF('Client Approval Tracker'!$F$7:$F$206,\"Sent\")+COUNTIF('Client Approval Tracker'!$F$7:$F$206,\"Revisions requested\")+COUNTIF('Client Approval Tracker'!$F$7:$F$206,\"On hold\")", '0'),
    ("Open vendor orders", "=COUNTIFS('Purchase Order - Order Tracker'!$A$7:$A$306,\"<>\",'Purchase Order - Order Tracker'!$J$7:$J$306,\"<>Delivered\",'Purchase Order - Order Tracker'!$J$7:$J$306,\"<>Cancelled\",'Purchase Order - Order Tracker'!$J$7:$J$306,\"<>Returned\")", '0'),
    ("Items not ordered", "=COUNTIFS('Master FF&E Schedule'!$D$7:$D$306,\"<>\",'Master FF&E Schedule'!$V$7:$V$306,\"\")", '0'),
    ("Delivered but not installed", "=COUNTIFS('Master FF&E Schedule'!$X$7:$X$306,\"Delivered\",'Master FF&E Schedule'!$Y$7:$Y$306,\"<>Installed\")", '0'),
]
# Six cards per row, three rows. Formulas are positioned to allow the Remaining Budget and Budget Used formulas to point at the card cells.
for idx, (label, formula, fmt) in enumerate(kpis):
    group = idx % 6
    row_group = idx // 6
    start_col = 1 + group * 3
    label_row = 5 + row_group * 4
    value_row = label_row + 1
    start_letter = get_column_letter(start_col)
    end_letter = get_column_letter(start_col + 2)
    dash.merge_cells(start_row=label_row, start_column=start_col, end_row=label_row, end_column=start_col + 2)
    dash.cell(label_row, start_col).value = label
    dash.cell(label_row, start_col).font = Font(size=9, bold=True, color=WHITE)
    dash.cell(label_row, start_col).fill = fill(TEAL)
    dash.cell(label_row, start_col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    dash.merge_cells(start_row=value_row, start_column=start_col, end_row=value_row + 1, end_column=start_col + 2)
    c = dash.cell(value_row, start_col)
    c.value = formula
    c.font = Font(size=17, bold=True, color=NAVY)
    c.fill = fill(PALE_BLUE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.number_format = fmt
    for rr in range(label_row, value_row + 2):
        for cc in range(start_col, start_col + 3):
            dash.cell(rr, cc).border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)
    dash.row_dimensions[label_row].height = 25
    dash.row_dimensions[value_row].height = 25
    dash.row_dimensions[value_row + 1].height = 20
# Fix Remaining budget and Budget used formulas to reference their actual merged-card top-left cells.
dash["J6"] = "=A6-G6"
dash["M6"] = "=IFERROR(G6/A6,0)"
# Conditional formatting for the dashboard values.
dash.conditional_formatting.add("J6", FormulaRule(formula=["=J6<0"], fill=fill(RED)))
dash.conditional_formatting.add("M6", FormulaRule(formula=["=M6>1"], fill=fill(RED)))
dash.conditional_formatting.add("P10", FormulaRule(formula=["=P10>0"], fill=fill(RED)))
dash.conditional_formatting.add("A14", FormulaRule(formula=["=A14>0"], fill=fill(YELLOW)))
dash.conditional_formatting.add("D14", FormulaRule(formula=["=D14>0"], fill=fill(YELLOW)))
dash.conditional_formatting.add("G14", FormulaRule(formula=["=G14>0"], fill=fill(YELLOW)))
dash.conditional_formatting.add("J14", FormulaRule(formula=["=J14>0"], fill=fill(RED)))
dash.conditional_formatting.add("A10", FormulaRule(formula=["=A10>0"], fill=fill(GREEN)))

# Dashboard chart area.
dash.merge_cells("A17:R17")
dash["A17"] = "VISUAL PROJECT OVERVIEW"
dash["A17"].font = Font(size=12, bold=True, color=WHITE)
dash["A17"].fill = fill(NAVY)
dash["A17"].alignment = Alignment(horizontal="left")

# Budget by room chart.
room_chart = BarChart()
room_chart.type = "bar"
room_chart.style = 10
room_chart.title = "Budget by Room"
room_chart.y_axis.title = "Room"
room_chart.x_axis.title = "Budget"
room_chart.height = 7.2
room_chart.width = 10.5
room_chart.legend = None
room_chart.add_data(Reference(helper, min_col=2, min_row=1, max_row=51), titles_from_data=True)
room_chart.set_categories(Reference(helper, min_col=1, min_row=2, max_row=51))
room_chart.varyColors = False
room_chart.series[0].graphicalProperties.solidFill = TEAL
room_chart.series[0].graphicalProperties.line.solidFill = TEAL
room_chart.dispBlanksAs = "gap"
dash.add_chart(room_chart, "A18")

# Budget versus actual.
bva_chart = BarChart()
bva_chart.type = "col"
bva_chart.style = 11
bva_chart.title = "Budget vs Actual"
bva_chart.y_axis.title = "Amount"
bva_chart.height = 7.2
bva_chart.width = 9.5
bva_chart.add_data(Reference(helper, min_col=5, min_row=1, max_row=4), titles_from_data=True)
bva_chart.set_categories(Reference(helper, min_col=4, min_row=2, max_row=4))
bva_chart.series[0].graphicalProperties.solidFill = GOLD
bva_chart.series[0].graphicalProperties.line.solidFill = GOLD
dash.add_chart(bva_chart, "K18")

# Item status breakdown.
item_chart = PieChart()
item_chart.title = "Item Status Breakdown"
item_chart.height = 7.2
item_chart.width = 9.5
item_chart.add_data(Reference(helper, min_col=8, min_row=1, max_row=1 + len(lists["OrderStatuses"])), titles_from_data=True)
item_chart.set_categories(Reference(helper, min_col=7, min_row=2, max_row=1 + len(lists["OrderStatuses"])))
item_chart.dataLabels = DataLabelList()
item_chart.dataLabels.showPercent = True
item_chart.dataLabels.showLeaderLines = True
dash.add_chart(item_chart, "K33")

# Delivery status breakdown.
delivery_chart = PieChart()
delivery_chart.title = "Delivery Status Breakdown"
delivery_chart.height = 7.2
delivery_chart.width = 9.5
delivery_chart.add_data(Reference(helper, min_col=11, min_row=1, max_row=1 + len(lists["DeliveryStatuses"])), titles_from_data=True)
delivery_chart.set_categories(Reference(helper, min_col=10, min_row=2, max_row=1 + len(lists["DeliveryStatuses"])))
delivery_chart.dataLabels = DataLabelList()
delivery_chart.dataLabels.showPercent = True
# no leader lines on second pie to keep it readable
dash.add_chart(delivery_chart, "A33")

# Category spend breakdown.
category_chart = BarChart()
category_chart.type = "bar"
category_chart.style = 12
category_chart.title = "Category Spend Breakdown"
category_chart.x_axis.title = "Final cost"
category_chart.y_axis.title = "Category"
category_chart.height = 8
category_chart.width = 18.5
category_chart.legend = None
category_chart.add_data(Reference(helper, min_col=14, min_row=1, max_row=1 + len(lists["FFECategories"])), titles_from_data=True)
category_chart.set_categories(Reference(helper, min_col=13, min_row=2, max_row=1 + len(lists["FFECategories"])))
category_chart.series[0].graphicalProperties.solidFill = TEAL
category_chart.series[0].graphicalProperties.line.solidFill = TEAL
dash.add_chart(category_chart, "A49")

dash.merge_cells("A66:R66")
dash["A66"] = "Dashboard data is formula-driven from the input tabs. If the dashboard is blank, start with the Client & Project Profile, Room List, Budget Tracker, and Master FF&E Schedule."
dash["A66"].font = Font(size=9, italic=True, color=TEAL)
dash["A66"].alignment = Alignment(wrap_text=True)
dash.freeze_panes = "A5"
dash.page_setup.orientation = "landscape"
dash.page_setup.fitToWidth = 1
dash.page_setup.fitToHeight = 0
dash.sheet_properties.pageSetUpPr.fitToPage = True

# ---------------------- Printable reports ----------------------
def make_report(title, subtitle, headers, widths, table_name, rows, formulas, tab_color=GOLD):
    sheet, end = make_data_sheet(title, subtitle, headers, widths, table_name, rows, formulas, tab_color)
    sheet["A4"] = "Printable report view • linked to the project tracker. Filter rows before printing or sharing."
    sheet["A4"].font = Font(size=9, italic=True, color=TEAL)
    return sheet, end

# Client-facing FF&E report linked to the master schedule.
report_headers = ["Room", "Category", "Item", "Description", "Brand / Vendor", "Product link", "Final cost", "Approval status"]
report_widths = [20, 19, 27, 34, 25, 34, 15, 18]
client_report, client_report_end = make_report("Client FF&E Report", "A clean, client-facing selection report. Filter by room or approval status before sharing.", report_headers, report_widths, "ClientFFEReport", 100, {
    1: lambda r: f'=IF(\'Master FF&E Schedule\'!D{r}="","",\'Master FF&E Schedule\'!B{r})',
    2: lambda r: f'=IF(\'Master FF&E Schedule\'!D{r}="","",\'Master FF&E Schedule\'!C{r})',
    3: lambda r: f'=IF(\'Master FF&E Schedule\'!D{r}="","",\'Master FF&E Schedule\'!D{r})',
    4: lambda r: f'=IF(\'Master FF&E Schedule\'!D{r}="","",\'Master FF&E Schedule\'!E{r})',
    5: lambda r: f'=IF(\'Master FF&E Schedule\'!D{r}="","",IF(\'Master FF&E Schedule\'!L{r}<>"",\'Master FF&E Schedule\'!L{r},\'Master FF&E Schedule\'!K{r}))',
    6: lambda r: f'=IF(\'Master FF&E Schedule\'!D{r}="","",\'Master FF&E Schedule\'!M{r})',
    7: lambda r: f'=IF(\'Master FF&E Schedule\'!D{r}="","",\'Master FF&E Schedule\'!T{r})',
    8: lambda r: f'=IF(\'Master FF&E Schedule\'!D{r}="","",\'Master FF&E Schedule\'!W{r})',
}, TEAL)
add_status_conditional_formatting(client_report, ["H7:H106"])

# Room budget report linked to room tracker.
room_report_headers = ["Room", "Budget allocation", "Estimated spend", "Actual spend", "Remaining budget", "Design status", "Priority", "Notes"]
room_report_widths = [24, 18, 18, 18, 19, 18, 14, 38]
room_report, room_report_end = make_report("Room-by-Room Budget Report", "Printable room-level budget report linked to the Room List - Area Tracker.", room_report_headers, room_report_widths, "RoomBudgetReport", 50, {
    1: lambda r: f'=\'Room List - Area Tracker\'!A{r}',
    2: lambda r: f'=\'Room List - Area Tracker\'!F{r}',
    3: lambda r: f'=\'Room List - Area Tracker\'!G{r}',
    4: lambda r: f'=\'Room List - Area Tracker\'!H{r}',
    5: lambda r: f'=\'Room List - Area Tracker\'!I{r}',
    6: lambda r: f'=\'Room List - Area Tracker\'!E{r}',
    7: lambda r: f'=\'Room List - Area Tracker\'!J{r}',
    8: lambda r: f'=\'Room List - Area Tracker\'!K{r}',
}, GOLD)
add_status_conditional_formatting(room_report, ["F7:F56", "G7:G56"])

# Vendor order report linked to order tracker.
vendor_report_headers = ["PO number", "Vendor", "Room", "Item", "Expected delivery", "Order status", "Tracking number", "Amount", "Paid status"]
vendor_report_widths = [15, 22, 20, 30, 21, 17, 25, 15, 16]
vendor_report, vendor_report_end = make_report("Vendor Order Report", "Printable open-order view for vendor follow-up and purchasing meetings.", vendor_report_headers, vendor_report_widths, "VendorOrderReport", 100, {
    1: lambda r: f'=IF(\'Purchase Order - Order Tracker\'!A{r}="","",\'Purchase Order - Order Tracker\'!A{r})',
    2: lambda r: f'=IF(\'Purchase Order - Order Tracker\'!A{r}="","",\'Purchase Order - Order Tracker\'!B{r})',
    3: lambda r: f'=IF(\'Purchase Order - Order Tracker\'!A{r}="","",\'Purchase Order - Order Tracker\'!C{r})',
    4: lambda r: f'=IF(\'Purchase Order - Order Tracker\'!A{r}="","",\'Purchase Order - Order Tracker\'!D{r})',
    5: lambda r: f'=IF(\'Purchase Order - Order Tracker\'!A{r}="","",\'Purchase Order - Order Tracker\'!H{r})',
    6: lambda r: f'=IF(\'Purchase Order - Order Tracker\'!A{r}="","",\'Purchase Order - Order Tracker\'!J{r})',
    7: lambda r: f'=IF(\'Purchase Order - Order Tracker\'!A{r}="","",\'Purchase Order - Order Tracker\'!K{r})',
    8: lambda r: f'=IF(\'Purchase Order - Order Tracker\'!A{r}="","",\'Purchase Order - Order Tracker\'!N{r})',
    9: lambda r: f'=IF(\'Purchase Order - Order Tracker\'!A{r}="","",\'Purchase Order - Order Tracker\'!O{r})',
}, TEAL)
add_status_conditional_formatting(vendor_report, ["F7:F106", "I7:I106"])

# Installation checklist linked to delivery tracker.
install_headers = ["Item", "Room", "Vendor", "Expected delivery", "Delivery status", "Damage check", "Install date", "Installer", "Installation status", "Punch list notes"]
install_widths = [28, 20, 22, 21, 17, 17, 15, 22, 19, 38]
install_report, install_report_end = make_report("Installation Checklist", "Printable receiving and installation checklist linked to Delivery & Installation Tracker.", install_headers, install_widths, "InstallationChecklist", 100, {
    1: lambda r: f'=IF(\'Delivery & Installation Tracker\'!A{r}="","",\'Delivery & Installation Tracker\'!A{r})',
    2: lambda r: f'=IF(\'Delivery & Installation Tracker\'!A{r}="","",\'Delivery & Installation Tracker\'!B{r})',
    3: lambda r: f'=IF(\'Delivery & Installation Tracker\'!A{r}="","",\'Delivery & Installation Tracker\'!C{r})',
    4: lambda r: f'=IF(\'Delivery & Installation Tracker\'!A{r}="","",\'Delivery & Installation Tracker\'!E{r})',
    5: lambda r: f'=IF(\'Delivery & Installation Tracker\'!A{r}="","",\'Delivery & Installation Tracker\'!G{r})',
    6: lambda r: f'=IF(\'Delivery & Installation Tracker\'!A{r}="","",\'Delivery & Installation Tracker\'!I{r})',
    7: lambda r: f'=IF(\'Delivery & Installation Tracker\'!A{r}="","",\'Delivery & Installation Tracker\'!L{r})',
    8: lambda r: f'=IF(\'Delivery & Installation Tracker\'!A{r}="","",\'Delivery & Installation Tracker\'!M{r})',
    9: lambda r: f'=IF(\'Delivery & Installation Tracker\'!A{r}="","",\'Delivery & Installation Tracker\'!N{r})',
    10: lambda r: f'=IF(\'Delivery & Installation Tracker\'!A{r}="","",\'Delivery & Installation Tracker\'!O{r})',
}, GOLD)
add_status_conditional_formatting(install_report, ["E7:E106", "F7:F106", "I7:I106"])

# Final punch-list report is intentionally editable, not linked: teams can create a closeout list quickly.
punch_headers = ["Punch-list item", "Room", "Due date", "Assigned to", "Status", "Notes"]
punch_widths = [34, 20, 15, 24, 18, 42]
punch, punch_end = make_report("Final Punch List", "Printable closeout checklist for final walkthrough, installation corrections, and handover.", punch_headers, punch_widths, "FinalPunchList", 100, {}, TEAL)
add_validation(punch, "RoomNames", "B7:B106")
add_validation(punch, "TaskStatuses", "E7:E106")
add_date_validation(punch, "C7:C106")
add_status_conditional_formatting(punch, ["E7:E106"])

# Add hyperlinks to the main navigation points.
for sheet in wb.worksheets:
    if sheet.title not in {"Welcome", "Lists & Helpers", "Dashboard Data"}:
        sheet["A3"].hyperlink = "#'Welcome'!A1"
        sheet["A3"].font = Font(name="Aptos", size=10, italic=True, color=TEAL, underline="single")

# Make all worksheets print neatly and consistently.
for sheet in wb.worksheets:
    if sheet.title not in {"Lists & Helpers", "Dashboard Data"}:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_margins.left = 0.25
        sheet.page_margins.right = 0.25
        sheet.page_margins.top = 0.5
        sheet.page_margins.bottom = 0.5
        sheet.page_margins.header = 0.2
        sheet.page_margins.footer = 0.2
        sheet.oddFooter.center.text = "&A"
        sheet.oddFooter.right.text = "Page &P of &N"
        sheet.oddFooter.center.size = 8
        sheet.oddFooter.right.size = 8

# Move the welcome and dashboard to the front explicitly; keep the helper sheet last.
order = ["Welcome", "Project Dashboard", "Client & Project Profile", "Room List - Area Tracker", "Master FF&E Schedule", "Furniture Schedule", "Fixtures Schedule", "Finishes Schedule", "Product Sourcing Tracker", "Vendor & Supplier Database", "Budget Tracker", "Purchase Order - Order Tracker", "Delivery & Installation Tracker", "Client Approval Tracker", "Project Timeline - Gantt", "Task & To-Do List", "Measurement Log", "Contractor - Installer Tracker", "Invoice & Payment Tracker", "Project Notes - Meeting Notes", "Client FF&E Report", "Room-by-Room Budget Report", "Vendor Order Report", "Installation Checklist", "Final Punch List", "Lists & Helpers", "Dashboard Data"]
wb._sheets = [wb[name] for name in order]

# Active view starts on the Welcome tab.
wb.active = 0
wb.save(OUT)

# Validate that the saved package opens and contains the intended tabs/formulas.
check = load_workbook(OUT, data_only=False)
assert check.sheetnames[:5] == ["Welcome", "Project Dashboard", "Client & Project Profile", "Room List - Area Tracker", "Master FF&E Schedule"]
assert check["Master FF&E Schedule"]["P7"].value == '=IF(D7="","",F7*O7)'
assert check["Master FF&E Schedule"]["T7"].value == '=IF(D7="","",P7+Q7+R7-S7)'
assert len(check["Project Dashboard"]._charts) == 5
print(f"Created {OUT} with {len(check.sheetnames)} sheets and {len(check['Project Dashboard']._charts)} dashboard charts.")
